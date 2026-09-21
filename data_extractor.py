"""
data_extractor.py - SEO Data Extractor & Multi-Tab Audit Workbook Generator
Extracts all crawl errors, creates an 'Index' summary sheet matching the client standard,
and generates individual tabs for each error category with detailed affected URLs.
"""

import io
import re
from urllib.parse import urlparse
from collections import Counter
import pandas as pd
try:
    from url_utils import is_pagination_url, get_base_unpaginated_url
except Exception:
    from seo_analyzer import is_pagination_url, get_base_unpaginated_url
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Excel tab name max length is 31 characters according to Excel specification
MAX_TAB_LEN = 31

def sanitize_sheet_title(title: str) -> str:
    """Sanitize sheet title for Excel: max 31 chars, no invalid characters [:\\/?*[]]."""
    clean = re.sub(r'[:\\/?*\[\]]', '', title).strip()
    if len(clean) > MAX_TAB_LEN:
        clean = clean[:MAX_TAB_LEN].strip()
    return clean or "Sheet"

def get_series(df: pd.DataFrame, col: str, default=""):
    """Safely get a Series from df, returning a default Series if col is absent."""
    if col in df.columns:
        return df[col]
    return pd.Series([default] * len(df), index=df.index)

def extract_all_seo_errors(analysis_result: dict) -> tuple[list[dict], dict[str, pd.DataFrame]]:
    """
    Extracts all errors and diagnostics from crawl results.
    Returns:
      - index_rows: List of dicts with keys: ['error_name', 'status', 'comments', 'sheet_name', 'count', 'is_error']
      - error_dfs: Dict mapping sheet_name -> DataFrame of affected URLs and diagnostics.
    """
    df_pages = analysis_result.get("df_pages", pd.DataFrame()).copy()
    df_links = analysis_result.get("df_links", pd.DataFrame()).copy()
    df_images = analysis_result.get("df_images", pd.DataFrame()).copy()
    
    index_rows = []
    error_dfs = {}

    if df_pages.empty:
        return index_rows, error_dfs

    # Helpers
    status_series = get_series(df_pages, "status_code", 0)
    indexable_series = get_series(df_pages, "is_indexable", True)
    is_indexable_200 = (status_series == 200) & (indexable_series == True)
    
    # Exclude noindex pages
    is_noindex_bool = get_series(df_pages, "is_noindex", False) == True
    meta_robots_noindex = get_series(df_pages, "meta_robots", "").fillna("").str.contains("noindex", case=False)
    is_noindex_mask = is_noindex_bool | meta_robots_noindex
    eval_pages = df_pages[is_indexable_200 & (~is_noindex_mask)]

    # -------------------------------------------------------------
    # 1. Internal links are broken (4xx / 5xx internal links)
    # -------------------------------------------------------------
    broken_internal_links = pd.DataFrame()
    if not df_links.empty and "url" in df_pages.columns:
        bad_urls = set(df_pages[get_series(df_pages, "status_code", 0) >= 400]["url"].dropna().unique())
        if bad_urls and "target_url" in df_links.columns:
            broken_internal_links = df_links[
                (get_series(df_links, "is_internal", False) == True) & 
                (df_links["target_url"].isin(bad_urls))
            ].copy()
    
    # Fallback to df_pages with 4xx and inlinks
    if broken_internal_links.empty:
        broken_pages = df_pages[get_series(df_pages, "status_code", 0) >= 400]
        if not broken_pages.empty:
            broken_count = len(broken_pages)
            b_tab = sanitize_sheet_title("Internal links are broken")
            cols = [c for c in ["url", "status_code", "status_description", "source_url", "anchor_text", "inlinks_count"] if c in broken_pages.columns]
            bdf = broken_pages[cols].copy().rename(columns={
                "url": "Broken Target URL",
                "status_code": "Status Code",
                "status_description": "HTTP Status",
                "source_url": "Source Page (Found On)",
                "anchor_text": "Anchor Text",
                "inlinks_count": "Inlinks Count"
            })
            bdf["Recommended Action"] = "Update internal hyperlink destination or implement a 301 redirect to a live page."
            error_dfs[b_tab] = bdf
            index_rows.append({
                "error_name": "Internal links are broken",
                "status": f"Failed ({broken_count})",
                "comments": f"Found {broken_count} internal URLs returning 4xx/5xx status codes.",
                "sheet_name": b_tab,
                "count": broken_count,
                "is_error": True
            })
        else:
            index_rows.append({
                "error_name": "Internal links are broken",
                "status": "Passed (0)",
                "comments": "All crawled internal links returned valid 200 OK responses.",
                "sheet_name": None,
                "count": 0,
                "is_error": False
            })
    else:
        broken_count = len(broken_internal_links)
        b_tab = sanitize_sheet_title("Internal links are broken")
        cols = [c for c in ["source_url", "target_url", "anchor_text"] if c in broken_internal_links.columns]
        bdf = broken_internal_links[cols].copy().rename(columns={
            "source_url": "Source Page (Found On)",
            "target_url": "Broken Target URL",
            "anchor_text": "Anchor Text"
        })
        bdf["Recommended Action"] = "Update or remove broken hyperlink on source page."
        error_dfs[b_tab] = bdf
        index_rows.append({
            "error_name": "Internal links are broken",
            "status": f"Failed ({broken_count})",
            "comments": f"Found {broken_count} broken internal hyperlinks pointing to dead URLs.",
            "sheet_name": b_tab,
            "count": broken_count,
            "is_error": True
        })

    # -------------------------------------------------------------
    # 3. Duplicate Meta Descriptions
    # -------------------------------------------------------------
    meta_s = get_series(eval_pages, "meta_description", "").fillna("").str.strip()
    dup_desc_df = pd.DataFrame()
    dup_descs = set()
    desc_counts = {}
    if not eval_pages.empty and "meta_description" in eval_pages.columns:
        valid_ep = eval_pages[meta_s != ""].copy()
        canon_col = get_series(valid_ep, "canonical_url", "")
        valid_ep["is_pagination"] = [is_pagination_url(u, c) for u, c in zip(valid_ep["url"], canon_col)]
        valid_ep["base_url"] = [get_base_unpaginated_url(u, c) for u, c in zip(valid_ep["url"], canon_col)]

        desc_to_bases = valid_ep.groupby("meta_description")["base_url"].apply(lambda s: set(s)).to_dict()
        dup_descs = {d for d, bases in desc_to_bases.items() if len(bases) > 1}
        # Only non-pagination pages count as duplicate errors
        dup_desc_df = valid_ep[valid_ep["meta_description"].isin(dup_descs) & (~valid_ep["is_pagination"])]
        desc_counts = {d: len(bases) for d, bases in desc_to_bases.items() if len(bases) > 1}
    
    dup_desc_count = len(dup_desc_df)
    dmd_tab = sanitize_sheet_title("Duplicate Meta Descriptions")
    if dup_desc_count > 0:
        cols = [c for c in ["url", "meta_description", "status_code"] if c in dup_desc_df.columns]
        ddf = dup_desc_df[cols].copy().rename(columns={
            "url": "Page URL",
            "meta_description": "Duplicate Meta Description",
            "status_code": "Status Code"
        })
        ddf["Duplicate Count"] = ddf["Duplicate Meta Description"].map(desc_counts)
        ddf["Recommended Action"] = "Write tailored, unique meta descriptions for each distinct page."
        error_dfs[dmd_tab] = ddf
        index_rows.append({
            "error_name": "Duplicate Meta Descriptions",
            "status": f"Failed ({dup_desc_count})",
            "comments": f"Found {dup_desc_count} pages sharing duplicate meta descriptions across {len(dup_descs)} groups.",
            "sheet_name": dmd_tab,
            "count": dup_desc_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Duplicate Meta Descriptions",
            "status": "Passed (0)",
            "comments": "All indexable pages have unique meta descriptions.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 4. 4XX status code
    # -------------------------------------------------------------
    sc_series = get_series(df_pages, "status_code", 0)
    c4xx_df = df_pages[(sc_series >= 400) & (sc_series < 500)]
    c4xx_count = len(c4xx_df)
    tab_4xx = sanitize_sheet_title("4XX status code")
    if c4xx_count > 0:
        cols = [c for c in ["url", "status_code", "status_description", "source_url", "anchor_text"] if c in c4xx_df.columns]
        edf = c4xx_df[cols].copy().rename(columns={
            "url": "Page URL",
            "status_code": "Status Code",
            "status_description": "Status Description",
            "source_url": "Discovered On",
            "anchor_text": "Anchor Text"
        })
        edf["Recommended Action"] = "Fix 404/403 URL, reinstate deleted content, or 301 redirect to relevant active page."
        error_dfs[tab_4xx] = edf
        index_rows.append({
            "error_name": "4XX status code",
            "status": f"Failed ({c4xx_count})",
            "comments": f"Found {c4xx_count} client-error URLs (404 Not Found, 403 Forbidden, etc.).",
            "sheet_name": tab_4xx,
            "count": c4xx_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "4XX status code",
            "status": "Passed (0)",
            "comments": "No 4XX client error status codes detected.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 5. Duplicate title tags
    # -------------------------------------------------------------
    title_s = get_series(eval_pages, "title", "").fillna("").str.strip()
    dup_title_df = pd.DataFrame()
    dup_titles = set()
    title_counts = {}
    if not eval_pages.empty and "title" in eval_pages.columns:
        valid_ep_t = eval_pages[title_s != ""].copy()
        canon_col_t = get_series(valid_ep_t, "canonical_url", "")
        valid_ep_t["is_pagination"] = [is_pagination_url(u, c) for u, c in zip(valid_ep_t["url"], canon_col_t)]
        valid_ep_t["base_url"] = [get_base_unpaginated_url(u, c) for u, c in zip(valid_ep_t["url"], canon_col_t)]

        title_to_bases = valid_ep_t.groupby("title")["base_url"].apply(lambda s: set(s)).to_dict()
        dup_titles = {t for t, bases in title_to_bases.items() if len(bases) > 1}
        # Only non-pagination pages count as duplicate errors
        dup_title_df = valid_ep_t[valid_ep_t["title"].isin(dup_titles) & (~valid_ep_t["is_pagination"])]
        title_counts = {t: len(bases) for t, bases in title_to_bases.items() if len(bases) > 1}
    
    dup_title_count = len(dup_title_df)
    dtt_tab = sanitize_sheet_title("Duplicate title tags")
    if dup_title_count > 0:
        cols = [c for c in ["url", "title", "status_code"] if c in dup_title_df.columns]
        tdf = dup_title_df[cols].copy().rename(columns={
            "url": "Page URL",
            "title": "Duplicate Title Tag",
            "status_code": "Status Code"
        })
        tdf["Duplicate Count"] = tdf["Duplicate Title Tag"].map(title_counts)
        tdf["Recommended Action"] = "Create distinct, keyword-focused title tags for each page."
        error_dfs[dtt_tab] = tdf
        index_rows.append({
            "error_name": "Duplicate title tags",
            "status": f"Failed ({dup_title_count})",
            "comments": f"Found {dup_title_count} pages sharing identical title tags across {len(dup_titles)} groups.",
            "sheet_name": dtt_tab,
            "count": dup_title_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Duplicate title tags",
            "status": "Passed (0)",
            "comments": "No duplicate title tags detected across indexable pages.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 5b. Missing title tags
    # -------------------------------------------------------------
    missing_title_mask = get_series(eval_pages, "title", "").fillna("").str.strip() == ""
    missing_title_df = eval_pages[missing_title_mask]
    missing_title_count = len(missing_title_df)
    mtt_tab = sanitize_sheet_title("Missing title tags")
    if missing_title_count > 0:
        cols = [c for c in ["url", "status_code"] if c in missing_title_df.columns]
        mtt_df = missing_title_df[cols].copy().rename(columns={"url": "Page URL", "status_code": "Status Code"})
        mtt_df["Recommended Action"] = "Add a descriptive, keyword-rich <title> tag between 30 and 60 characters."
        error_dfs[mtt_tab] = mtt_df
        index_rows.append({
            "error_name": "Missing title tags",
            "status": f"Failed ({missing_title_count})",
            "comments": f"Found {missing_title_count} indexable pages completely missing a <title> tag.",
            "sheet_name": mtt_tab,
            "count": missing_title_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Missing title tags",
            "status": "Passed (0)",
            "comments": "All indexable pages have a title tag.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 5c. Title tags over 60 Chars (>600px)
    # -------------------------------------------------------------
    t_len_s = get_series(eval_pages, "title_length", 0)
    t_px_s = get_series(eval_pages, "title_pixel_width", 0)
    long_t_mask = (t_len_s > 60) | (t_px_s > 600)
    long_title_df = eval_pages[long_t_mask]
    long_title_count = len(long_title_df)
    ltt_tab = sanitize_sheet_title("Titles over 60 Chars")
    if long_title_count > 0:
        cols = [c for c in ["url", "title", "title_length", "title_pixel_width", "status_code"] if c in long_title_df.columns]
        ltdf = long_title_df[cols].copy().rename(columns={
            "url": "Page URL",
            "title": "Page Title",
            "title_length": "Length (Chars)",
            "title_pixel_width": "Pixel Width (px)",
            "status_code": "Status Code"
        })
        ltdf["Recommended Action"] = "Shorten title to under 60 characters (<600px) to avoid ellipsis truncation in search results."
        error_dfs[ltt_tab] = ltdf
        index_rows.append({
            "error_name": "Title tags over 60 Chars (>600px)",
            "status": f"Attention Needed ({long_title_count})",
            "comments": f"Found {long_title_count} pages with titles exceeding SERP display limits (>60 chars or >600px).",
            "sheet_name": ltt_tab,
            "count": long_title_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Title tags over 60 Chars (>600px)",
            "status": "Passed (0)",
            "comments": "All page titles fit comfortably within Google SERP pixel limits.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 6. Minify JavaScript and CSS files
    # -------------------------------------------------------------
    unminified_list = []
    # Check evaluated crawled pages for unminified assets
    if not df_pages.empty:
        for _, pr in df_pages.iterrows():
            page_u = pr.get("url", "")
            assets = pr.get("unminified_assets", [])
            if isinstance(assets, list) and assets:
                for asset in assets:
                    clean_a = str(asset).split("?")[0].lower()
                    unminified_list.append({
                        "Page URL (Found On)": page_u,
                        "Unminified Resource URL": str(asset),
                        "Resource Type": "JavaScript (.js)" if clean_a.endswith(".js") else "Stylesheet (.css)",
                        "Recommended Action": "Minify and bundle JS/CSS files using build tools (e.g. Terser, CSSNano, esbuild) or CDN auto-minification to reduce transfer payload and boost Core Web Vitals (FCP, LCP)."
                    })
    # Also check df_links as supplementary
    if not df_links.empty:
        for _, lr in df_links.iterrows():
            tgt = str(lr.get("target_url", ""))
            src = str(lr.get("source_url", ""))
            clean_tgt = tgt.split("?")[0].lower()
            if (clean_tgt.endswith(".js") or clean_tgt.endswith(".css")) and ".min." not in clean_tgt and not clean_tgt.endswith(".min.js") and not clean_tgt.endswith(".min.css"):
                unminified_list.append({
                    "Page URL (Found On)": src,
                    "Unminified Resource URL": tgt,
                    "Resource Type": "JavaScript (.js)" if clean_tgt.endswith(".js") else "Stylesheet (.css)",
                    "Recommended Action": "Minify and bundle JS/CSS files using build tools (e.g. Terser, CSSNano, esbuild) or CDN auto-minification to reduce transfer payload and boost Core Web Vitals (FCP, LCP)."
                })
    
    unmin_tab = sanitize_sheet_title("Minify JS & CSS")
    if unminified_list:
        unmin_df = pd.DataFrame(unminified_list).drop_duplicates(subset=["Page URL (Found On)", "Unminified Resource URL"]).head(500)
        unmin_count = len(unmin_df)
        error_dfs[unmin_tab] = unmin_df
        index_rows.append({
            "error_name": "Minify JavaScript and CSS files",
            "status": f"Attention Needed ({unmin_count})",
            "comments": f"Found {unmin_count} unminified script (.js) and stylesheet (.css) resources that delay page rendering.",
            "sheet_name": unmin_tab,
            "count": unmin_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Minify JavaScript and CSS files",
            "status": "Passed (0)",
            "comments": "All detected JavaScript and CSS assets appear minified or optimized.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 7. Underscores in the URL
    # -------------------------------------------------------------
    def has_underscore(u: str) -> bool:
        path = urlparse(str(u)).path
        return "_" in path

    underscore_df = pd.DataFrame()
    if "url" in df_pages.columns and not df_pages.empty:
        underscore_df = df_pages[df_pages["url"].apply(has_underscore)]
    under_count = len(underscore_df)
    under_tab = sanitize_sheet_title("Underscores in the URL")
    if under_count > 0:
        cols = [c for c in ["url", "status_code"] if c in underscore_df.columns]
        udf = underscore_df[cols].copy().rename(columns={"url": "Page URL", "status_code": "Status Code"})
        udf["Recommended Clean Slug"] = udf["Page URL"].apply(
            lambda u: urlparse(u)._replace(path=urlparse(u).path.replace("_", "-")).geturl()
        )
        udf["Recommended Action"] = "Replace underscores with hyphens in URL slugs per Google SEO guidelines and 301 redirect old URLs."
        error_dfs[under_tab] = udf
        index_rows.append({
            "error_name": "Underscores in the URL",
            "status": f"Attention Needed ({under_count})",
            "comments": f"Found {under_count} URLs with underscores '_' in the slug path.",
            "sheet_name": under_tab,
            "count": under_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Underscores in the URL",
            "status": "Passed (0)",
            "comments": "All URLs use hyphens or clean slugs without underscores.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 8. Don't have meta descriptions
    # -------------------------------------------------------------
    no_desc_mask = get_series(eval_pages, "meta_description", "").fillna("").str.strip() == ""
    no_desc_df = eval_pages[no_desc_mask]
    no_desc_count = len(no_desc_df)
    dhm_tab = sanitize_sheet_title("Don't have meta descriptions")
    if no_desc_count > 0:
        cols = [c for c in ["url", "title", "status_code"] if c in no_desc_df.columns]
        mdf = no_desc_df[cols].copy().rename(columns={
            "url": "Page URL",
            "title": "Page Title",
            "status_code": "Status Code"
        })
        mdf["Recommended Action"] = "Add an engaging, keyword-rich meta description between 120-155 characters."
        error_dfs[dhm_tab] = mdf
        index_rows.append({
            "error_name": "Don't have meta descriptions",
            "status": f"Failed ({no_desc_count})",
            "comments": f"Found {no_desc_count} indexable pages completely missing meta descriptions.",
            "sheet_name": dhm_tab,
            "count": no_desc_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Don't have meta descriptions",
            "status": "Passed (0)",
            "comments": "All indexable crawled pages have meta descriptions.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 8b. Meta descriptions over 160 Chars
    # -------------------------------------------------------------
    d_len_s = get_series(eval_pages, "meta_description_length", 0)
    long_desc_mask = d_len_s > 160
    long_desc_df = eval_pages[long_desc_mask]
    long_desc_count = len(long_desc_df)
    ld_tab = sanitize_sheet_title("Desc over 160 Chars")
    if long_desc_count > 0:
        cols = [c for c in ["url", "meta_description", "meta_description_length", "status_code"] if c in long_desc_df.columns]
        lddf = long_desc_df[cols].copy().rename(columns={
            "url": "Page URL",
            "meta_description": "Meta Description",
            "meta_description_length": "Length (Chars)",
            "status_code": "Status Code"
        })
        lddf["Recommended Action"] = "Shorten meta description to 120-155 characters to avoid SERP ellipsis truncation."
        error_dfs[ld_tab] = lddf
        index_rows.append({
            "error_name": "Meta descriptions over 160 Chars",
            "status": f"Attention Needed ({long_desc_count})",
            "comments": f"Found {long_desc_count} pages with meta descriptions over 160 characters.",
            "sheet_name": ld_tab,
            "count": long_desc_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Meta descriptions over 160 Chars",
            "status": "Passed (0)",
            "comments": "No meta descriptions exceed standard SERP snippet limits.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 9. Missing H1
    # -------------------------------------------------------------
    h1_text = get_series(eval_pages, "h1", "").fillna("").str.strip()
    h1_count_ser = get_series(eval_pages, "h1_count", 0).fillna(0)
    missing_h1_mask = (h1_text == "") | (h1_count_ser == 0)
    missing_h1_df = eval_pages[missing_h1_mask]
    missing_h1_count = len(missing_h1_df)
    mh1_tab = sanitize_sheet_title("Missing H1")
    if missing_h1_count > 0:
        cols = [c for c in ["url", "title", "status_code"] if c in missing_h1_df.columns]
        h1df = missing_h1_df[cols].copy().rename(columns={
            "url": "Page URL",
            "title": "Page Title",
            "status_code": "Status Code"
        })
        h1df["Recommended Action"] = "Add exactly one primary H1 heading describing the core topic of the page."
        error_dfs[mh1_tab] = h1df
        index_rows.append({
            "error_name": "Missing H1",
            "status": f"Failed ({missing_h1_count})",
            "comments": f"Found {missing_h1_count} indexable pages without a primary H1 tag.",
            "sheet_name": mh1_tab,
            "count": missing_h1_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Missing H1",
            "status": "Passed (0)",
            "comments": "All indexable pages contain at least one valid H1 heading tag.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 9b. Multiple H1 tags
    # -------------------------------------------------------------
    multi_h1_mask = h1_count_ser > 1
    multi_h1_df = eval_pages[multi_h1_mask]
    multi_h1_count = len(multi_h1_df)
    multi_h1_tab = sanitize_sheet_title("Multiple H1 tags")
    if multi_h1_count > 0:
        cols = [c for c in ["url", "h1_count", "h1", "h1_2", "status_code"] if c in multi_h1_df.columns]
        mhdf = multi_h1_df[cols].copy().rename(columns={
            "url": "Page URL",
            "h1_count": "Total H1 Count",
            "h1": "First H1 Tag",
            "h1_2": "Second H1 Tag",
            "status_code": "Status Code"
        })
        mhdf["Recommended Action"] = "Remove superfluous H1 tags so each page has exactly one primary H1 heading."
        error_dfs[multi_h1_tab] = mhdf
        index_rows.append({
            "error_name": "Multiple H1 tags",
            "status": f"Failed ({multi_h1_count})",
            "comments": f"Found {multi_h1_count} indexable pages containing multiple H1 tags (more than 1 H1).",
            "sheet_name": multi_h1_tab,
            "count": multi_h1_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Multiple H1 tags",
            "status": "Passed (0)",
            "comments": "No multiple H1 tags found; all pages have at most one primary H1.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 9c. Duplicate H1 tags
    # -------------------------------------------------------------
    valid_h1_ep = eval_pages[h1_text != ""].copy()
    dup_h1_df = pd.DataFrame()
    dup_h1_count = 0
    if not valid_h1_ep.empty:
        canon_col_h = get_series(valid_h1_ep, "canonical_url", "")
        valid_h1_ep["is_pagination"] = [is_pagination_url(u, c) for u, c in zip(valid_h1_ep["url"], canon_col_h)]
        valid_h1_ep["base_url"] = [get_base_unpaginated_url(u, c) for u, c in zip(valid_h1_ep["url"], canon_col_h)]
        h1_to_bases = valid_h1_ep.groupby("h1")["base_url"].apply(lambda s: set(s)).to_dict()
        dup_h1_set = {h for h, bases in h1_to_bases.items() if len(bases) > 1}
        dup_h1_df = valid_h1_ep[valid_h1_ep["h1"].isin(dup_h1_set) & (~valid_h1_ep["is_pagination"])]
        dup_h1_count = len(dup_h1_df)

    dh1_tab = sanitize_sheet_title("Duplicate H1 tags")
    if dup_h1_count > 0:
        cols = [c for c in ["url", "h1", "status_code"] if c in dup_h1_df.columns]
        dhdf = dup_h1_df[cols].copy().rename(columns={
            "url": "Page URL",
            "h1": "Duplicate H1 Tag",
            "status_code": "Status Code"
        })
        dhdf["Recommended Action"] = "Provide unique H1 headings for distinct landing pages."
        error_dfs[dh1_tab] = dhdf
        index_rows.append({
            "error_name": "Duplicate H1 tags",
            "status": f"Failed ({dup_h1_count})",
            "comments": f"Found {dup_h1_count} pages sharing duplicate H1 headings across distinct pages.",
            "sheet_name": dh1_tab,
            "count": dup_h1_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Duplicate H1 tags",
            "status": "Passed (0)",
            "comments": "All indexable pages have unique H1 headings.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 10. Orphaned pages in sitemaps / site
    # -------------------------------------------------------------
    orphan_mask = get_series(eval_pages, "is_orphan", False) == True
    orphan_df = eval_pages[orphan_mask]
    orphan_count = len(orphan_df)
    orphan_tab = sanitize_sheet_title("Orphaned pages in sitemaps")
    if orphan_count > 0:
        cols = [c for c in ["url", "title", "status_code", "inlinks_count"] if c in orphan_df.columns]
        odf = orphan_df[cols].copy().rename(columns={
            "url": "Page URL",
            "title": "Page Title",
            "status_code": "Status Code",
            "inlinks_count": "Inlinks (0)"
        })
        odf["Recommended Action"] = "Add internal hyperlinks pointing to this URL from top navigation, category pages, or relevant blogs."
        error_dfs[orphan_tab] = odf
        index_rows.append({
            "error_name": "Orphaned pages in sitemaps",
            "status": f"Attention Needed ({orphan_count})",
            "comments": f"Found {orphan_count} indexable orphan pages with 0 internal inlinks.",
            "sheet_name": orphan_tab,
            "count": orphan_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Orphaned pages in sitemaps",
            "status": "Passed (0)",
            "comments": "No orphan pages found; all pages receive internal incoming links.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 10b. Canonical URL points to alternative URL
    # -------------------------------------------------------------
    canon_stat_s = get_series(eval_pages, "canonical_status", "")
    canon_alt_mask = canon_stat_s == "Canonicalised"
    canon_alt_df = eval_pages[canon_alt_mask]
    canon_alt_count = len(canon_alt_df)
    can_tab = sanitize_sheet_title("Canonical URL issues")
    if canon_alt_count > 0:
        cols = [c for c in ["url", "canonical_url", "status_code"] if c in canon_alt_df.columns]
        candf = canon_alt_df[cols].copy().rename(columns={
            "url": "Crawled Page URL",
            "canonical_url": "Points to Canonical URL",
            "status_code": "Status Code"
        })
        candf["Recommended Action"] = "Verify that this canonical target is the desired master version to consolidate link equity."
        error_dfs[can_tab] = candf
        index_rows.append({
            "error_name": "Canonical URL points to alternative URL",
            "status": f"Attention Needed ({canon_alt_count})",
            "comments": f"Found {canon_alt_count} pages where canonical tag points to an alternative URL.",
            "sheet_name": can_tab,
            "count": canon_alt_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Canonical URL points to alternative URL",
            "status": "Passed (0)",
            "comments": "All canonical tags are self-referential or clean.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 11. Images over 100 KB
    # -------------------------------------------------------------
    heavy_images = df_images[get_series(df_images, "is_over_100kb", False) == True] if not df_images.empty else pd.DataFrame()
    heavy_count = len(heavy_images)
    img_tab = sanitize_sheet_title("Images over 100 KB")
    if heavy_count > 0:
        cols = [c for c in ["page_url", "image_url", "size_kb", "content_type"] if c in heavy_images.columns]
        ihdf = heavy_images[cols].copy().rename(columns={
            "page_url": "Page URL (Found On)",
            "image_url": "Heavy Image URL",
            "size_kb": "File Size (KB)",
            "content_type": "Format"
        }).head(1000)
        ihdf["Recommended Action"] = "Compress or convert image to modern WebP format under 100 KB."
        error_dfs[img_tab] = ihdf
        index_rows.append({
            "error_name": "Images over 100 KB",
            "status": f"Failed ({heavy_count})",
            "comments": f"Found {heavy_count} heavy image files exceeding 100 KB.",
            "sheet_name": img_tab,
            "count": heavy_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Images over 100 KB",
            "status": "Passed (0)",
            "comments": "All images are optimized under 100 KB.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 11b. Slow server response time (>1.5s)
    # -------------------------------------------------------------
    slow_pages = df_pages[(get_series(df_pages, "response_time_ms", 0) > 1500) | (get_series(df_pages, "latency_ms", 0) > 1500)]
    slow_count = len(slow_pages)
    slow_tab = sanitize_sheet_title("Slow response time")
    if slow_count > 0:
        cols = [c for c in ["url", "status_code", "latency_ms"] if c in slow_pages.columns]
        sdf = slow_pages[cols].copy().rename(columns={
            "url": "Slow Page URL",
            "status_code": "Status Code",
            "latency_ms": "Latency (ms)"
        })
        sdf["Recommended Action"] = "Optimize database queries, enable server-side caching or CDN."
        error_dfs[slow_tab] = sdf
        index_rows.append({
            "error_name": "Slow server response time (>1.5s)",
            "status": f"Attention Needed ({slow_count})",
            "comments": f"Found {slow_count} pages taking more than 1.5 seconds to respond.",
            "sheet_name": slow_tab,
            "count": slow_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Slow server response time (>1.5s)",
            "status": "Passed (0)",
            "comments": "All crawled pages responded within optimal latency (<1.5s).",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 12. Missing Image Alt Text
    # -------------------------------------------------------------
    missing_alt = pd.DataFrame()
    if not df_images.empty:
        missing_alt = df_images[get_series(df_images, "has_alt", True) == False]
    alt_count = len(missing_alt)
    alt_tab = sanitize_sheet_title("Images Missing Alt Text")
    if alt_count > 0:
        cols = [c for c in ["page_url", "image_url", "loading"] if c in missing_alt.columns]
        altdf = missing_alt[cols].copy().rename(columns={
            "page_url": "Page URL",
            "image_url": "Image URL",
            "loading": "Loading Attribute"
        }).head(1000)
        altdf["Recommended Action"] = "Add descriptive alt attributes to assist accessibility and Google Image ranking."
        error_dfs[alt_tab] = altdf
        index_rows.append({
            "error_name": "Images Missing Alt Text",
            "status": f"Failed ({alt_count})",
            "comments": f"Found {alt_count} images lacking descriptive alt attributes.",
            "sheet_name": alt_tab,
            "count": alt_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Images Missing Alt Text",
            "status": "Passed (0)",
            "comments": "All detected images include descriptive alt attributes.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    # -------------------------------------------------------------
    # 13. Redirect Chains & Loops (3XX)
    # -------------------------------------------------------------
    chain_mask = (get_series(df_pages, "is_redirect_chain", False) == True) | (get_series(df_pages, "is_redirect_loop", False) == True)
    chain_df = df_pages[chain_mask]
    chain_count = len(chain_df)
    chain_tab = sanitize_sheet_title("Redirect Chains & Loops")
    if chain_count > 0:
        cols = [c for c in ["url", "status_code", "redirect_chain_str", "redirect_hops", "final_url"] if c in chain_df.columns]
        cdf = chain_df[cols].copy().rename(columns={
            "url": "Start URL",
            "status_code": "Initial Status",
            "redirect_chain_str": "Redirect Chain",
            "redirect_hops": "Hops Count",
            "final_url": "Final Destination URL"
        })
        cdf["Recommended Action"] = "Update inlinks to point directly to the final 200 OK target URL."
        error_dfs[chain_tab] = cdf
        index_rows.append({
            "error_name": "Redirect Chains & Loops",
            "status": f"Failed ({chain_count})",
            "comments": f"Found {chain_count} URLs involved in multiple redirect hops or loops.",
            "sheet_name": chain_tab,
            "count": chain_count,
            "is_error": True
        })
    else:
        index_rows.append({
            "error_name": "Redirect Chains & Loops",
            "status": "Passed (0)",
            "comments": "No multi-hop redirect chains or infinite loops detected.",
            "sheet_name": None,
            "count": 0,
            "is_error": False
        })

    return index_rows, error_dfs

def build_error_audit_excel_workbook(index_rows: list[dict], error_dfs: dict[str, pd.DataFrame]) -> bytes:
    """
    Builds a beautifully formatted Excel workbook matching the screenshot specification:
    - Sheet 1: 'Index' with merged green header, columns [Errors, Status, Comments]
    - Following Sheets: Dedicated tab for each detected error containing affected URLs.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    font_main_header = Font(name="Calibri", size=13, bold=True, color="000000")
    font_col_header = Font(name="Calibri", size=11, bold=True, color="000000")
    font_data = Font(name="Calibri", size=10, color="1F2937")
    font_link = Font(name="Calibri", size=10, color="0B57D0", underline="single")
    font_status_fail = Font(name="Calibri", size=10, bold=True, color="991B1B")
    font_status_pass = Font(name="Calibri", size=10, bold=True, color="166534")

    fill_main_green = PatternFill(start_color="B6D7A8", end_color="B6D7A8", fill_type="solid") # Google Sheets Pale Green
    fill_col_green = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")  # Lighter pale green
    fill_fail = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_ai_header = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid") # Indigo AI Header
    font_ai_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_ai_cell = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid") # Soft Indigo Tint

    border_thin = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # -------------------------------------------------------------
    # CREATE 'Index' SHEET
    # -------------------------------------------------------------
    ws_index = wb.create_sheet(title="Index")
    ws_index.views.sheetView[0].showGridLines = True

    has_ai_plan = any("suggested_action_plan" in item or "ai_action_plan" in item for item in index_rows)
    merge_range = "A1:D1" if has_ai_plan else "A1:C1"
    banner_cols = ["A", "B", "C", "D"] if has_ai_plan else ["A", "B", "C"]

    # Row 1: Merged 'Index'
    ws_index.merge_cells(merge_range)
    cell_idx = ws_index["A1"]
    cell_idx.value = "Index"
    cell_idx.font = font_main_header
    cell_idx.fill = fill_main_green
    cell_idx.alignment = align_center
    ws_index.row_dimensions[1].height = 28

    # Apply borders & fill across merged cells
    for col_l in banner_cols:
        c = ws_index[f"{col_l}1"]
        c.border = border_thin
        c.fill = fill_main_green

    # Row 2: Header Columns
    headers = ["Errors", "Status", "Comments", "Suggested Action Plan & Developer Guide"] if has_ai_plan else ["Errors", "Status", "Comments"]
    ws_index.row_dimensions[2].height = 22
    for col_num, h_text in enumerate(headers, 1):
        c = ws_index.cell(row=2, column=col_num)
        c.value = h_text
        if any(k in h_text for k in ["Suggested", "Developer", "AI "]):
            c.font = font_ai_header
            c.fill = fill_ai_header
        else:
            c.font = font_col_header
            c.fill = fill_col_green
        c.alignment = align_center if col_num == 2 else align_left
        c.border = border_thin

    # Rows 3+: Data Rows
    current_row = 3
    for item in index_rows:
        ws_index.row_dimensions[current_row].height = 20
        c_err = ws_index.cell(row=current_row, column=1)
        c_stat = ws_index.cell(row=current_row, column=2)
        c_comm = ws_index.cell(row=current_row, column=3)

        c_err.value = item["error_name"]
        sheet_target = item.get("sheet_name")
        # If there's an associated error tab, add hyperlink!
        if sheet_target and sheet_target in error_dfs:
            c_err.hyperlink = f"#'{sheet_target}'!A1"
            c_err.font = font_link
        else:
            c_err.font = font_data

        c_err.alignment = align_left
        c_err.border = border_thin

        c_stat.value = item["status"]
        if item.get("is_error", False):
            c_stat.fill = fill_fail
            c_stat.font = font_status_fail
        else:
            c_stat.fill = fill_pass
            c_stat.font = font_status_pass
        c_stat.alignment = align_center
        c_stat.border = border_thin

        c_comm.value = item["comments"]
        c_comm.font = font_data
        c_comm.alignment = align_left
        c_comm.border = border_thin

        if has_ai_plan:
            c_ai = ws_index.cell(row=current_row, column=4)
            c_ai.value = item.get("suggested_action_plan", item.get("ai_action_plan", ""))
            c_ai.font = font_data
            c_ai.fill = fill_ai_cell
            c_ai.alignment = align_left
            c_ai.border = border_thin

        current_row += 1

    # Column widths for Index sheet
    ws_index.column_dimensions["A"].width = 38
    ws_index.column_dimensions["B"].width = 24
    ws_index.column_dimensions["C"].width = 65
    if has_ai_plan:
        ws_index.column_dimensions["D"].width = 75
    ws_index.freeze_panes = "A3"

    # -------------------------------------------------------------
    # CREATE INDIVIDUAL ERROR SHEETS
    # -------------------------------------------------------------
    for sheet_name, df_err in error_dfs.items():
        if df_err.empty:
            continue

        ws_err = wb.create_sheet(title=sheet_name)
        ws_err.views.sheetView[0].showGridLines = True

        # Header Row
        ws_err.row_dimensions[1].height = 24
        col_names = list(df_err.columns)
        ai_col_indices = set()
        for col_idx, col_name in enumerate(col_names, 1):
            c = ws_err.cell(row=1, column=col_idx)
            c.value = str(col_name)
            if any(k in str(col_name) for k in ["Suggested", "Developer Guide", "Action Plan", "AI "]):
                c.font = font_ai_header
                c.fill = fill_ai_header
                ai_col_indices.add(col_idx)
            else:
                c.font = font_col_header
                c.fill = fill_col_green
            c.alignment = align_center if "Status" in col_name or "Count" in col_name else align_left
            c.border = border_thin

        # Data Rows
        for r_idx, row_data in enumerate(df_err.itertuples(index=False), 2):
            ws_err.row_dimensions[r_idx].height = 20
            for c_idx, val in enumerate(row_data, 1):
                c = ws_err.cell(row=r_idx, column=c_idx)
                c.value = val if val is not None else ""
                c.font = font_data
                c.border = border_thin
                c.alignment = align_left
                if c_idx in ai_col_indices:
                    c.fill = fill_ai_cell
                # Link formatting for URLs
                if isinstance(val, str) and val.startswith("http"):
                    c.font = font_link

        # Auto-adjust column widths
        for col in ws_err.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_err.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 70)

        ws_err.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

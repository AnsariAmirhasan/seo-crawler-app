"""
Advance GSC Data Analysis Module for CrawlPilot SEO Suite
Integrates deep Google Search Console (GSC) export analysis, striking-distance quick wins,
low-hanging fruit CTR optimization, content gaps, questions intent, traffic projection modeling,
target page mapping engine with brand homepage recognition, AI title/H1/meta-description generation,
interactive page override, and complete Excel export.
"""

import io
import re
import os
import json
import zipfile
import urllib.parse
import numpy as np
import pandas as pd
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
import matplotlib.pyplot as plt
from wordcloud import WordCloud
from typing import Dict, Any, Tuple, Optional, List

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def read_csv_robust(file_bytes_or_buffer) -> pd.DataFrame:
    """Robustly reads a CSV file trying different encodings, delimiters, and strip BOM."""
    if isinstance(file_bytes_or_buffer, bytes):
        raw_bytes = file_bytes_or_buffer
    else:
        raw_bytes = file_bytes_or_buffer.read()

    # Try common encodings and separators
    for enc in ['utf-8-sig', 'utf-8', 'latin1', 'cp1252', 'utf-16']:
        for sep in [None, ',', ';', '\t']:
            try:
                if sep is None:
                    df = pd.read_csv(io.BytesIO(raw_bytes), encoding=enc, sep=None, engine='python')
                else:
                    df = pd.read_csv(io.BytesIO(raw_bytes), encoding=enc, sep=sep)

                if len(df.columns) >= 2 or (len(df) > 0 and len(df.columns) == 1):
                    # Clean column names (strip BOM, quotes, spaces)
                    cleaned_cols = [str(c).replace('\ufeff', '').strip().strip('"').strip("'") for c in df.columns]
                    # If columns were merged because sep sniffing didn't catch it, continue trying
                    if len(cleaned_cols) == 1 and (';' in cleaned_cols[0] or ',' in cleaned_cols[0] or '\t' in cleaned_cols[0]):
                        continue
                    df.columns = cleaned_cols
                    return df
            except Exception:
                continue

    # Ultimate fallback default
    try:
        df = pd.read_csv(io.BytesIO(raw_bytes), encoding_errors='replace')
        df.columns = [str(c).replace('\ufeff', '').strip().strip('"').strip("'") for c in df.columns]
        return df
    except Exception as e:
        st.warning(f"Warning parsing CSV: {e}")
        return pd.DataFrame()


def extract_csv_from_zip(zip_file) -> Dict[str, pd.DataFrame]:
    """Extract all CSV files from the uploaded ZIP file with case-insensitive matching."""
    csv_data = {}
    with zipfile.ZipFile(zip_file, 'r') as zip_ref:
        for file in zip_ref.namelist():
            if file.lower().endswith('.csv') and not file.startswith('__MACOSX'):
                clean_name = os.path.basename(file).replace('.csv', '').replace('.CSV', '')
                try:
                    df = read_csv_robust(zip_ref.read(file))
                    if not df.empty:
                        csv_data[clean_name] = df
                        csv_data[clean_name.lower()] = df
                except Exception as e:
                    st.warning(f"Warning: Could not parse {file}: {e}")
    return csv_data


def clean_ctr_column(df: pd.DataFrame, column_name: str = 'ctr') -> pd.DataFrame:
    """Convert CTR from percentage string ('2.5%') or scaled float to a proper decimal float."""
    for col in df.columns:
        if col.lower() == column_name.lower() or 'ctr' in col.lower() or 'klickrate' in col.lower() or 'taux' in col.lower():
            try:
                if df[col].dtype == object:
                    df[col] = (
                        df[col]
                        .astype(str)
                        .str.replace('%', '', regex=False)
                        .str.replace(',', '.', regex=False)
                        .str.strip()
                    )
                    df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0.0)
                else:
                    df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0.0)

                if df[col].max() > 1.0:
                    df[col] = df[col] / 100.0
            except Exception:
                df[col] = 0.0
            break
    return df


def standardize_columns(df: pd.DataFrame, default_entity_type: str = 'query') -> pd.DataFrame:
    """
    Standardize column names to lowercase with underscores and resolve GSC naming variations.
    Guarantees that query/page/country, clicks, impressions, ctr, position columns exist.
    """
    if df.empty:
        return df

    new_cols = []
    for col in df.columns:
        c = str(col).replace('\ufeff', '').strip().lower().replace(' ', '_').replace('-', '_')
        # Check CTR FIRST so 'click_through_rate' does NOT get matched to 'clicks'
        if 'ctr' in c or 'click_through' in c or 'clickthrough' in c or 'klickrate' in c or 'taux' in c or '%' in c:
            new_cols.append('ctr')
        elif any(k in c for k in ['query', 'queries', 'keyword', 'consulta', 'suchanfrag', 'requêt']):
            new_cols.append('query')
        elif any(k in c for k in ['page', 'url', 'landing_page', 'seite', 'página']):
            new_cols.append('page')
        elif any(k in c for k in ['countr', 'land', 'pays', 'país']):
            new_cols.append('country')
        elif any(k in c for k in ['clicks', 'click', 'clics', 'clic', 'klicks', 'klick', 'cliques', 'clique']):
            new_cols.append('clicks')
        elif any(k in c for k in ['impress', 'view', 'vista', 'vue', 'impressionen']):
            new_cols.append('impressions')
        elif any(k in c for k in ['position', 'pos', 'rang', 'rank', 'platz']):
            new_cols.append('position')
        else:
            new_cols.append(c)

    df.columns = new_cols

    # Positional mapping fallback if standard names weren't present
    if default_entity_type not in df.columns and len(df.columns) > 0:
        df.rename(columns={df.columns[0]: default_entity_type}, inplace=True)

    if 'clicks' not in df.columns:
        if len(df.columns) > 1 and df.columns[1] not in ['query', 'page', 'country']:
            df.rename(columns={df.columns[1]: 'clicks'}, inplace=True)
        else:
            df['clicks'] = 0

    if 'impressions' not in df.columns:
        if len(df.columns) > 2 and df.columns[2] not in ['query', 'page', 'country', 'clicks']:
            df.rename(columns={df.columns[2]: 'impressions'}, inplace=True)
        else:
            df['impressions'] = 0

    if 'ctr' not in df.columns:
        if len(df.columns) > 3 and df.columns[3] not in ['query', 'page', 'country', 'clicks', 'impressions']:
            df.rename(columns={df.columns[3]: 'ctr'}, inplace=True)
        else:
            df['ctr'] = 0.0

    if 'position' not in df.columns:
        if len(df.columns) > 4 and df.columns[4] not in ['query', 'page', 'country', 'clicks', 'impressions', 'ctr']:
            df.rename(columns={df.columns[4]: 'position'}, inplace=True)
        else:
            df['position'] = 0.0

    return df


def sanitize_dataframe(df: pd.DataFrame, default_entity_type: str = 'query') -> pd.DataFrame:
    """Standardizes, deduplicates, cleans CTR, and casts numeric columns safely without throwing TypeError."""
    if df.empty:
        return df

    # Standardize column headers
    df = standardize_columns(df, default_entity_type=default_entity_type)

    # Deduplicate columns if any duplicate name exists
    df = df.loc[:, ~df.columns.duplicated(keep='first')].copy()

    # Clean CTR column
    df = clean_ctr_column(df, 'ctr')

    # Convert numeric columns safely (handling Series or 1D DataFrame)
    for col in ['clicks', 'impressions']:
        if col in df.columns:
            series = df[col].iloc[:, 0] if isinstance(df[col], pd.DataFrame) else df[col]
            df[col] = pd.to_numeric(series, errors='coerce').fillna(0).astype(int)
        else:
            df[col] = 0

    for col in ['position', 'ctr']:
        if col in df.columns:
            series = df[col].iloc[:, 0] if isinstance(df[col], pd.DataFrame) else df[col]
            df[col] = pd.to_numeric(series, errors='coerce').fillna(0.0).astype(float)
        else:
            df[col] = 0.0

    # Mathematical guarantee for CTR: prevents display bugs (e.g. 0.00% when clicks > 0)
    if 'impressions' in df.columns and 'clicks' in df.columns:
        df['ctr'] = np.where(df['impressions'] > 0, df['clicks'] / df['impressions'], 0.0)

    return df


def extract_base_domain(pages_df: pd.DataFrame, default: str = "https://example.com") -> str:
    """Extracts base website protocol + domain from pages dataframe or session crawl results."""
    crawl_res = st.session_state.get("crawl_results")
    if crawl_res and crawl_res.get("start_url"):
        parsed = urllib.parse.urlparse(crawl_res.get("start_url"))
        if parsed.netloc:
            return f"{parsed.scheme or 'https'}://{parsed.netloc}"

    if not pages_df.empty and 'page' in pages_df.columns:
        for url in pages_df['page'].dropna():
            u = str(url).strip()
            if u.startswith('http://') or u.startswith('https://'):
                parsed = urllib.parse.urlparse(u)
                if parsed.netloc:
                    return f"{parsed.scheme}://{parsed.netloc}"

    return default


def extract_brand_tokens(domain_str: str) -> set:
    """Extracts brand root name and tokens from the domain (e.g. airaindia.com -> {'aira', 'india', 'airaindia'})."""
    parsed = urllib.parse.urlparse(domain_str if domain_str.startswith('http') else 'https://' + domain_str)
    netloc = parsed.netloc.lower()
    netloc = re.sub(r'^www\.', '', netloc)
    parts = netloc.split('.')
    main_name = parts[0] if parts else netloc

    brand_tokens = {main_name}
    for suffix in ['india', 'global', 'valves', 'valve', 'tech', 'corp', 'ltd', 'pvt', 'automation']:
        if main_name.endswith(suffix) and len(main_name) > len(suffix) + 2:
            root = main_name[:-len(suffix)]
            brand_tokens.add(root)
            brand_tokens.add(suffix)

    return brand_tokens


def is_brand_query(query: str, brand_tokens: set) -> bool:
    """Determines whether a query is a brand navigational query that should target the Homepage."""
    q_words = re.findall(r'[a-zA-Z0-9]+', str(query).lower())
    if not q_words:
        return False

    has_brand = any(w in brand_tokens for w in q_words)
    if not has_brand:
        return False

    generic_brand_modifiers = {
        'euro', 'india', 'valves', 'valve', 'automation', 'pvt', 'ltd',
        'company', 'corp', 'group', 'products', 'manufacturer', 'supplier',
        'official', 'website', 'contact', 'address', 'near', 'me', 'online'
    }
    non_brand_words = [w for w in q_words if w not in brand_tokens and w not in generic_brand_modifiers]

    # If query is purely brand + company modifiers or very short (1-2 words), it's a brand query
    if len(non_brand_words) == 0 or len(q_words) <= 2:
        return True

    return False


def tokenize_text(text: str) -> set:
    """Tokenize and stem keywords for URL slug matching."""
    words = re.findall(r'[a-zA-Z0-9]+', str(text).lower())
    stems = set()
    for w in words:
        stems.add(w)
        if w.endswith('es') and len(w) > 4:
            stems.add(w[:-2])
            stems.add(w[:-1])
        elif w.endswith('s') and len(w) > 3:
            stems.add(w[:-1])
        if w.endswith('ing') and len(w) > 5:
            stems.add(w[:-3])
        if w.endswith('ed') and len(w) > 4:
            stems.add(w[:-2])
    return stems


def map_queries_to_pages(queries_df: pd.DataFrame, pages_df: pd.DataFrame, base_domain: str = "") -> pd.DataFrame:
    """
    Intelligently maps each query to an existing URL from pages_df (or crawled pages)
    with strict Homepage matching for brand queries and anti-blog hijacking for commercial queries.
    """
    df_q = queries_df.copy()
    if df_q.empty or 'query' not in df_q.columns:
        return df_q

    if not base_domain:
        base_domain = extract_base_domain(pages_df)
    base_domain = base_domain.rstrip('/')
    homepage_url = base_domain + '/'

    brand_tokens = extract_brand_tokens(base_domain)

    # Candidate URLs from pages_df and crawl results
    candidate_urls = []
    if not pages_df.empty and 'page' in pages_df.columns:
        candidate_urls.extend([str(u).strip() for u in pages_df['page'].dropna().unique() if str(u).strip()])

    crawl_data = st.session_state.get("crawl_results")
    if crawl_data and isinstance(crawl_data, dict):
        crawled_list = crawl_data.get("pages", [])
        if isinstance(crawled_list, list):
            for cp in crawled_list:
                if isinstance(cp, dict) and cp.get("url"):
                    candidate_urls.append(cp["url"].strip())

    candidate_urls = list(dict.fromkeys(candidate_urls))

    target_pages = []
    target_types = []

    common_industry_generics = {
        'valve', 'valves', 'type', 'types', 'price', 'pricing', 'spec', 'specs',
        'guide', 'online', 'india', 'buy', 'best', 'manufacturers', 'suppliers',
        'in', 'on', 'at', 'to', 'for', 'of', 'with', 'by', 'a', 'an', 'the', 'and', 'or', 'is', 'are'
    }

    for q in df_q['query']:
        q_clean = str(q).strip().lower()

        # 1. BRAND QUERIES MUST TARGET HOMEPAGE
        if is_brand_query(q_clean, brand_tokens):
            target_pages.append(homepage_url)
            target_types.append("🏠 Homepage (Brand)")
            continue

        # 2. Extract core keywords for non-brand queries
        q_words = re.findall(r'[a-zA-Z0-9]+', q_clean)
        core_q_words = [w for w in q_words if w not in common_industry_generics]
        if not core_q_words:
            core_q_words = q_words

        slug_clean = re.sub(r'[^a-z0-9]+', '-', q_clean).strip('-')
        best_score = 0.0
        best_url = ""

        # 3. Check candidate URLs
        for u in candidate_urls:
            p = urllib.parse.urlparse(u).path.lower()
            if p in ['', '/']:
                continue

            # Exact slug match gets highest priority
            if slug_clean and slug_clean in p:
                best_score = 2.0
                best_url = u
                break

            path_words = set(re.findall(r'[a-zA-Z0-9]+', p))
            matched_core = [w for w in core_q_words if w in path_words]

            # Core keyword MUST be present in path to prevent false matches (e.g. nrv vs butterfly)
            if len(matched_core) == 0:
                continue

            score = len(matched_core) / len(core_q_words)

            # Penalize blog posts for commercial/product queries so product pages win
            if any(b in p for b in ['/blog/', '/article/', '/post/', '/news/']):
                score *= 0.25
            else:
                score *= 1.25

            if score > best_score:
                best_score = score
                best_url = u

        if best_score >= 0.50 and best_url:
            target_pages.append(best_url)
            target_types.append("Existing Page")
        else:
            is_question = any(q_clean.startswith(w) for w in ['what', 'how', 'why', 'when', 'where', 'which', 'who', 'can', 'does', 'is', 'vs'])
            prefix = "blog" if is_question else "products"
            rec_url = f"{base_domain}/{prefix}/{slug_clean}"
            target_pages.append(rec_url)
            target_types.append("💡 Recommended: Create New Page")

    df_q['target_page'] = target_pages
    df_q['target_type'] = target_types
    return df_q


def call_ai_api(api_key: str, provider: str, model: str, system_prompt: str, user_prompt: str) -> str:
    """Call Google Gemini, OpenAI, or Anthropic Claude API for GSC SEO suggestions."""
    clean_model = model.split(" ")[0].strip() if " (" in model else model.strip()

    if provider == "Google Gemini":
        from google import genai
        client = genai.Client(api_key=api_key)
        try:
            response = client.models.generate_content(
                model=clean_model,
                contents=f"{system_prompt}\n\nUSER PROMPT:\n{user_prompt}"
            )
            return response.text or ""
        except Exception as e:
            err = str(e).lower()
            if "404" in err or "not found" in err or "not_found" in err:
                for fallback in ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-1.5-flash"]:
                    try:
                        response = client.models.generate_content(
                            model=fallback,
                            contents=f"{system_prompt}\n\nUSER PROMPT:\n{user_prompt}"
                        )
                        return response.text or ""
                    except Exception:
                        continue
            raise e

    elif provider == "ChatGPT (OpenAI)":
        from openai import OpenAI
        client = OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model=clean_model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            temperature=0.7
        )
        return response.choices[0].message.content or ""

    elif provider == "Claude (Anthropic)":
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model=clean_model,
            max_tokens=4000,
            system=system_prompt,
            messages=[{"role": "user", "content": user_prompt}]
        )
        return response.content[0].text if response.content else ""

    return ""


def generate_algorithmic_seo_suggestions(
    query: str,
    target_page: str,
    target_type: str,
    position: float,
    ctr: float,
    category: str,
    base_domain: str = "",
    brand_tokens: set = None
) -> Dict[str, str]:
    """Generates deterministic, high-converting SEO titles, H1s, meta descriptions, and action plans."""
    q_title = query.title()
    brand_tokens = brand_tokens or set()

    # Brand query handling
    if is_brand_query(query, brand_tokens) or "Homepage" in target_type:
        title = f"{q_title} | Official Website & Valve Solutions [2026]"
        if len(title) > 60: title = f"{q_title} | Official Manufacturer & Exporter"
        h1 = f"{q_title}: Industrial Valves & Automation Solutions"
        meta = f"Welcome to the official {q_title} website. Discover high-performance industrial valves, automated actuators, and technical specifications. Contact us today!"
        action = f"Brand navigational search ranking at Pos {position:.1f}. Target Homepage ({target_page}). Ensure brand name is in the Homepage Title Tag, H1, and add Organization Schema with logo and social links."
        return {
            "query": query,
            "target_page": target_page,
            "target_type": "🏠 Homepage (Brand)",
            "ai_title": title,
            "ai_h1": h1,
            "ai_meta_description": meta,
            "action_plan": action
        }

    if category == "Low-Hanging Fruit":
        title = f"{q_title} - Price, Specs & Buying Guide [2026]"
        if len(title) > 60: title = f"{q_title} [2026 Price & Specs]"
        h1 = f"{q_title}: Industrial Specifications & Technical Guide"
        meta = f"Looking for top-grade {query}? Explore verified technical specs, competitive prices, and fast dispatch. Request a quick quote today!"
        action = f"Prime SERP rank (Pos {position:.1f}) but low CTR ({ctr*100:.1f}%). Revise title tag on {target_page} with brackets, pricing hints, and trust triggers to double click-through rate."
    elif category == "Quick Win":
        title = f"Best {q_title}: Top Models, Ratings & Reviews"
        if len(title) > 60: title = f"{q_title} - Models & Reviews [2026]"
        h1 = f"Complete Guide to {q_title}: Applications & Benefits"
        meta = f"Discover everything about {query}. Learn operational benefits, technical standards, and expert maintenance tips. Read the complete breakdown."
        action = f"Keyword ranks on Page 2 (Pos {position:.1f}). Add a dedicated H2 section on '{query}' on {target_page}, insert 2 internal links from top-traffic pages, and expand body copy by 300 words."
    elif category == "Question":
        title = f"{q_title}? Step-by-Step Answer & Guide"
        if len(title) > 60: title = f"{q_title}? Answer & Guide"
        h1 = f"{q_title}: Complete Technical Explanation"
        meta = f"Need an answer to: {query}? Here is the complete step-by-step technical explanation and expert checklist. Find out more inside!"
        action = f"Featured Snippet opportunity! Add an H2 with an exact 40-50 word direct answer block followed by bulleted steps, and apply FAQ Schema."
    elif category == "Content Gap":
        title = f"{q_title} Solutions - Manufacturer & Supplier"
        if len(title) > 60: title = f"{q_title} Manufacturer & Supply"
        h1 = f"{q_title}: Industrial Equipment Solutions"
        meta = f"High-durability {query} engineered for demanding applications. Meets global quality standards. Contact our technical team for custom specs."
        action = f"Keyword ranks at Pos {position:.1f}. Target this term with a dedicated landing page ({target_page}), complete product catalog, and downloadable spec sheets."
    else:
        title = f"{q_title} | Premium Solutions & Catalog"
        if len(title) > 60: title = f"{q_title} | Expert Solutions"
        h1 = f"{q_title}: Industrial Overview"
        meta = f"Explore leading {query} options with proven performance and reliability. View specifications and request assistance from our specialists."
        action = f"Top-performing query driving organic volume. Maintain rankings with periodic content refreshes and fresh internal links."

    return {
        "query": query,
        "target_page": target_page,
        "target_type": target_type,
        "ai_title": title,
        "ai_h1": h1,
        "ai_meta_description": meta,
        "action_plan": action
    }


def generate_ai_seo_action_plan(
    items_to_process: List[Dict[str, Any]],
    api_key: str = "",
    provider: str = "Google Gemini",
    model: str = "gemini-3.8-pro-preview",
    candidate_pages: List[str] = None,
    base_domain: str = "",
    brand_tokens: set = None
) -> List[Dict[str, Any]]:
    """Generates complete SEO suggestions using Gemini/OpenAI/Claude with full site pages context and guaranteed fallback."""
    results = []
    candidate_pages = candidate_pages or []
    brand_tokens = brand_tokens or set()
    homepage_url = base_domain.rstrip('/') + '/' if base_domain else ""

    # If no API key provided, generate high-quality programmatic recommendations instantly
    if not api_key:
        for it in items_to_process:
            sug = generate_algorithmic_seo_suggestions(
                query=it['query'],
                target_page=it.get('target_page', ''),
                target_type=it.get('target_type', 'Existing Page'),
                position=float(it.get('position', 0)),
                ctr=float(it.get('ctr', 0)),
                category=it.get('category', 'Top Query'),
                base_domain=base_domain,
                brand_tokens=brand_tokens
            )
            it.update(sug)
            results.append(it)
        return results

    # Format sample of candidate pages for the AI prompt
    clean_pages_sample = [homepage_url] if homepage_url else []
    for cp in candidate_pages[:35]:
        if cp not in clean_pages_sample:
            clean_pages_sample.append(cp)

    pages_prompt_text = "\n".join([f"- {url}" for url in clean_pages_sample[:30]])

    batch_size = 12
    for i in range(0, len(items_to_process), batch_size):
        batch = items_to_process[i:i+batch_size]

        system_prompt = (
            "You are a World-Class Enterprise SEO Strategist, Information Architect & Content Optimization Expert.\n"
            f"WEBSITE DOMAIN: {base_domain}\n"
            f"BRAND NAMES & TERMS: {', '.join(brand_tokens)}\n\n"
            f"ACTUAL PAGES EXISTING ON THIS WEBSITE:\n{pages_prompt_text}\n\n"
            "STRICT TARGET PAGE & METADATA RULES:\n"
            "1. BRAND QUERIES: If a query is for the company brand (e.g. 'aira', 'aira euro', 'aira valves', 'company name'), "
            f"the target_page MUST ALWAYS be the Homepage ('{homepage_url}'). Brand queries MUST NOT be assigned to blog posts!\n"
            "2. COMMERCIAL / PRODUCT QUERIES: Pick the most specific product/solution page from the ACTUAL PAGES list above. "
            f"DO NOT assign commercial queries to a blog post if a product page exists. If no existing page exists, recommend a clean URL like '{base_domain}/products/{{slug}}'.\n"
            "3. INFORMATIONAL / QUESTIONS: Assign to an existing guide/blog or recommend '{base_domain}/blog/{{slug}}'.\n"
            "4. Return STRICTLY a valid JSON array of objects. Do not include markdown if possible, or standard ```json.\n"
            "Each object must have these exact keys:\n"
            "- query: the input query string\n"
            "- target_page: the best matching or recommended target URL\n"
            "- target_type: '🏠 Homepage (Brand)', 'Existing Page', or '💡 Recommended: Create New Page'\n"
            "- ai_title: CTR-optimized page title (50-60 characters max, include numbers, brackets, power words)\n"
            "- ai_h1: clear, authoritative H1 tag\n"
            "- ai_meta_description: compelling meta description (130-155 characters max with high-intent CTA)\n"
            "- action_plan: specific actionable on-page advice (e.g. title tweaks, internal linking, schema markup, content expansion)"
        )

        user_prompt = "Generate SEO optimization recommendations for these Google Search Console queries:\n"
        for idx, row in enumerate(batch):
            user_prompt += (
                f"{idx+1}. Query: '{row['query']}' | Category: '{row.get('category', 'General')}' | "
                f"Position: {row.get('position', 0):.1f} | CTR: {row.get('ctr', 0)*100:.1f}%\n"
            )

        try:
            raw_text = call_ai_api(api_key, provider, model, system_prompt, user_prompt)
            clean_json = re.sub(r'^```json\s*', '', raw_text.strip(), flags=re.MULTILINE)
            clean_json = re.sub(r'```$', '', clean_json.strip(), flags=re.MULTILINE).strip()
            parsed_data = json.loads(clean_json)

            lookup = {str(item.get('query', '')).strip().lower(): item for item in parsed_data if isinstance(item, dict)}
            for row in batch:
                q_key = str(row['query']).strip().lower()
                ai_item = lookup.get(q_key)
                if ai_item:
                    ai_page = ai_item.get('target_page', '')
                    if ai_page:
                        row['target_page'] = ai_page
                        if ai_page.rstrip('/') == base_domain.rstrip('/'):
                            row['target_type'] = '🏠 Homepage (Brand)'
                        elif any(ai_page == cp for cp in candidate_pages):
                            row['target_type'] = 'Existing Page'
                        else:
                            row['target_type'] = ai_item.get('target_type', '💡 Recommended: Create New Page')
                    row['ai_title'] = ai_item.get('ai_title', '')
                    row['ai_h1'] = ai_item.get('ai_h1', '')
                    row['ai_meta_description'] = ai_item.get('ai_meta_description', '')
                    row['action_plan'] = ai_item.get('action_plan', '')
                    results.append(row)
                else:
                    sug = generate_algorithmic_seo_suggestions(
                        query=row['query'],
                        target_page=row.get('target_page', ''),
                        target_type=row.get('target_type', 'Existing Page'),
                        position=float(row.get('position', 0)),
                        ctr=float(row.get('ctr', 0)),
                        category=row.get('category', 'Top Query'),
                        base_domain=base_domain,
                        brand_tokens=brand_tokens
                    )
                    row.update(sug)
                    results.append(row)
        except Exception:
            # Fallback for this batch on any API failure
            for row in batch:
                sug = generate_algorithmic_seo_suggestions(
                    query=row['query'],
                    target_page=row.get('target_page', ''),
                    target_type=row.get('target_type', 'Existing Page'),
                    position=float(row.get('position', 0)),
                    ctr=float(row.get('ctr', 0)),
                    category=row.get('category', 'Top Query'),
                    base_domain=base_domain,
                    brand_tokens=brand_tokens
                )
                row.update(sug)
                results.append(row)

    return results


def export_action_plan_to_excel(df: pd.DataFrame) -> bytes:
    """Exports SEO Action Plan into a beautifully formatted Excel file with colored headers and auto column widths."""
    output = io.BytesIO()

    if not OPENPYXL_AVAILABLE:
        with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
            df.to_excel(writer, index=False, sheet_name='SEO Action Plan')
        return output.getvalue()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='SEO Action Plan')
        workbook = writer.book
        worksheet = writer.sheets['SEO Action Plan']

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in range(2, len(df) + 2):
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        for col_idx, col_name in enumerate(df.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = max([len(str(val or '')) for val in df[col_name].dropna()] + [len(col_name)])
            col_width = min(max(max_len + 3, 12), 45)
            worksheet.column_dimensions[col_letter].width = col_width

    return output.getvalue()


def generate_sample_gsc_data() -> Dict[str, pd.DataFrame]:
    """Generates a realistic sample GSC dataset for quick testing and demonstration."""
    sample_queries = [
        ("aira", 107, 5818, 0.0184, 3.4),
        ("aira euro", 96, 628, 0.1529, 1.1),
        ("aira valves", 88, 445, 0.1978, 1.2),
        ("wafer type butterfly valve", 0, 434, 0.0000, 1.1),
        ("knife edge gate valve", 2, 379, 0.0053, 1.1),
        ("best industrial valves", 1420, 18500, 0.0768, 4.2),
        ("electric actuator valve manufacturer", 980, 14200, 0.0690, 3.8),
        ("knife gate valve installation guide", 450, 12600, 0.0357, 12.4),
        ("pneumatic actuator vs electric actuator", 320, 11400, 0.0281, 14.2),
        ("how does a knife gate valve work", 290, 9800, 0.0296, 11.8),
        ("what is an industrial valve actuator", 260, 8900, 0.0292, 13.5),
        ("ansi 150 valve pressure ratings", 180, 15400, 0.0117, 18.2),
        ("high temperature butterfly valve specs", 140, 16800, 0.0083, 21.4),
        ("zero leakage valve maintenance", 110, 13200, 0.0083, 19.5),
        ("cast steel vs stainless steel valves", 95, 7800, 0.0122, 15.3),
        ("how to calibrate electric actuator", 85, 6400, 0.0133, 16.1),
        ("valve actuator wiring diagram", 70, 5200, 0.0135, 17.8),
        ("why do valve seats leak under pressure", 60, 4800, 0.0125, 14.9),
        ("valve torque calculation formula", 55, 18200, 0.0030, 68.4),
        ("industrial piping flange dimensions table", 42, 22400, 0.0019, 74.2),
        ("severe service slurry gate valves", 38, 14500, 0.0026, 62.1),
        ("top valve suppliers in gujarat", 520, 2400, 0.2167, 1.8),
        ("industrial automation equipment", 410, 28000, 0.0146, 2.8),
        ("flanged ball valve dimensions", 380, 19500, 0.0195, 2.9),
        ("where to buy actuated valves in india", 190, 4200, 0.0452, 8.4),
        ("when to replace valve gland packing", 115, 3900, 0.0295, 9.6),
    ]
    df_q = pd.DataFrame(sample_queries, columns=['query', 'clicks', 'impressions', 'ctr', 'position'])

    sample_pages = [
        ("https://www.airaindia.com/", 2450, 32400, 0.0756, 1.4),
        ("https://www.airaindia.com/products/", 1120, 18200, 0.0615, 3.2),
        ("https://www.airaindia.com/double-off-set-disc-wafer-type-butterfly-valve/", 420, 8900, 0.0472, 1.8),
        ("https://www.airaindia.com/knife-edge-gate-valve/", 380, 7400, 0.0514, 1.6),
        ("https://www.airaindia.com/blog/aira-make-piston-valve/", 95, 3200, 0.0297, 8.4),
        ("https://www.airaindia.com/blog/best-control-valves-manufactured-by-aira-euro/", 60, 2800, 0.0214, 9.1),
        ("https://www.airaindia.com/blog/aira-valves-in-textile-industry/", 45, 1900, 0.0237, 10.4),
        ("https://www.airaindia.com/technical-specs/valve-torque-calculator", 85, 34500, 0.0025, 65.4),
        ("https://www.airaindia.com/about-us", 650, 4200, 0.1548, 1.4),
        ("https://www.airaindia.com/contact-us", 420, 3100, 0.1355, 1.6),
    ]
    df_p = pd.DataFrame(sample_pages, columns=['page', 'clicks', 'impressions', 'ctr', 'position'])

    sample_countries = [
        ("India", 3450, 42500, 0.0812, 5.2),
        ("United States", 1240, 32100, 0.0386, 9.8),
        ("United Arab Emirates", 820, 14800, 0.0554, 7.1),
        ("United Kingdom", 540, 12400, 0.0435, 11.2),
        ("Germany", 410, 9600, 0.0427, 12.5),
        ("Saudi Arabia", 380, 8900, 0.0427, 8.9),
    ]
    df_c = pd.DataFrame(sample_countries, columns=['country', 'clicks', 'impressions', 'ctr', 'position'])

    return {
        'Queries': df_q,
        'queries': df_q,
        'Pages': df_p,
        'pages': df_p,
        'Countries': df_c,
        'countries': df_c
    }


def render_advance_gsc_analysis_page():
    """Main Streamlit renderer for the Advance GSC data Analysis service."""

    # 1. Header Banner
    st.markdown("""
    <div style="background: radial-gradient(130% 120% at 50% -10%, #1e1b4b 0%, #0F172A 60%, #020617 100%); padding: 2.5rem 2rem 2rem; border-radius: 20px; border: 1px solid rgba(139, 92, 246, 0.35); text-align: center; margin-bottom: 2rem; box-shadow: 0 20px 45px -10px rgba(0,0,0,0.6);">
        <div style="display: inline-flex; align-items: center; gap: 8px; background: rgba(139, 92, 246, 0.15); border: 1px solid rgba(139, 92, 246, 0.4); padding: 5px 16px; border-radius: 999px; margin-bottom: 12px;">
            <span style="font-size: 0.85rem; font-weight: 800; color: #C4B5FD; letter-spacing: 0.06em; text-transform: uppercase;">📈 Google Search Console Deep Intelligence</span>
        </div>
        <h1 style="font-size: 2.75rem; font-weight: 900; color: #FFFFFF; letter-spacing: -0.03em; margin: 0 0 0.6rem 0; line-height: 1.15;">
            Advance GSC Data Analysis
        </h1>
        <p style="color: #94A3B8; font-size: 1.05rem; max-width: 780px; margin: 0 auto; line-height: 1.6;">
            Deep-dive into your Google Search Console performance data. Map every query to accurate target URLs (Brand ➔ Homepage), optimize CTR, and generate AI-powered Title, H1, Meta Descriptions, and complete action plan exports.
        </p>
    </div>
    """, unsafe_allow_html=True)

    # 2. Upload and Data Acquisition Area
    col_upload, col_sample = st.columns([3.5, 1.5])
    with col_upload:
        uploaded_file = st.file_uploader(
            "📤 Upload Google Search Console Export (ZIP or CSV format):",
            type=['zip', 'csv'],
            help="In Google Search Console, click 'Export' ➔ 'Download ZIP' (or upload Queries.csv directly)."
        )
    with col_sample:
        st.markdown("<div style='margin-top: 28px;'></div>", unsafe_allow_html=True)
        load_sample = st.button("🧪 Load Sample GSC Data", use_container_width=True, help="Load an industrial demo dataset to preview all reports and features instantly.")

    csv_files = {}

    if uploaded_file is not None:
        try:
            with st.spinner("📦 Extracting and parsing Google Search Console data..."):
                if uploaded_file.name.lower().endswith('.zip'):
                    csv_files = extract_csv_from_zip(uploaded_file)
                else:
                    single_df = read_csv_robust(uploaded_file.getvalue())
                    clean_name = os.path.splitext(uploaded_file.name)[0]
                    csv_files[clean_name] = single_df
                    csv_files[clean_name.lower()] = single_df
                    csv_files['queries'] = single_df
                    csv_files['Queries'] = single_df

                st.session_state["gsc_cached_data"] = csv_files
                st.toast("✅ GSC Data successfully loaded and processed!", icon="🚀")
        except Exception as e:
            st.error(f"❌ Error processing file: {str(e)}")
            st.info("💡 Please make sure the file is directly exported from Google Search Console.")
            return
    elif load_sample or st.session_state.get("gsc_use_sample"):
        st.session_state["gsc_use_sample"] = True
        csv_files = generate_sample_gsc_data()
        st.session_state["gsc_cached_data"] = csv_files
        st.info("ℹ️ Currently displaying **Sample Industrial GSC Demo Dataset**. Upload your real export above anytime to analyze your live site.")
    elif "gsc_cached_data" in st.session_state:
        csv_files = st.session_state["gsc_cached_data"]
    else:
        # Empty State Helper Card
        st.markdown("""
        <div style="background: rgba(15, 23, 42, 0.7); border: 1px dashed rgba(139, 92, 246, 0.4); border-radius: 16px; padding: 2.5rem 2rem; text-align: center; margin-top: 1rem;">
            <div style="font-size: 2.8rem; margin-bottom: 0.8rem;">📥</div>
            <h3 style="color: #F8FAFC; margin-bottom: 0.5rem; font-weight: 800;">No GSC Data Loaded Yet</h3>
            <p style="color: #94A3B8; font-size: 0.95rem; max-width: 580px; margin: 0 auto 1.5rem; line-height: 1.6;">
                Export your Google Search Console performance data (Performance ➔ Export ➔ Download ZIP), then upload it above or click <b>'Load Sample GSC Data'</b> to see a full live demonstration.
            </p>
            <div style="display: flex; justify-content: center; gap: 15px; font-size: 0.84rem; color: #CBD5E1;">
                <span>✔ <b>Queries.csv</b> Analysis</span>
                <span>✔ <b>Pages.csv</b> Analysis</span>
                <span>✔ <b>Countries.csv</b> Analysis</span>
                <span>✔ Target Page Mapping</span>
                <span>✔ AI Meta & Complete File Export</span>
            </div>
        </div>
        """, unsafe_allow_html=True)
        return

    queries_df = pd.DataFrame()
    for k in csv_files:
        if any(w in k.lower() for w in ['quer', 'suchanfrag', 'consult', 'requêt', 'keyword']):
            queries_df = csv_files[k].copy()
            break

    pages_df = pd.DataFrame()
    for k in csv_files:
        if any(w in k.lower() for w in ['page', 'seite', 'págin', 'url']):
            pages_df = csv_files[k].copy()
            break

    countries_df = pd.DataFrame()
    for k in csv_files:
        if any(w in k.lower() for w in ['countr', 'land', 'pays', 'país']):
            countries_df = csv_files[k].copy()
            break

    if queries_df.empty and len(csv_files) > 0:
        first_key = list(csv_files.keys())[0]
        queries_df = csv_files[first_key].copy()

    if queries_df.empty:
        st.error("❌ 'Queries.csv' was not found inside the uploaded ZIP. Please ensure your Google Search Console export contains Queries data.")
        return

    queries_df = sanitize_dataframe(queries_df, default_entity_type='query')

    if not pages_df.empty:
        pages_df = sanitize_dataframe(pages_df, default_entity_type='page')

    if not countries_df.empty:
        countries_df = sanitize_dataframe(countries_df, default_entity_type='country')

    # Detect or configure website base domain for target page generation
    auto_domain = extract_base_domain(pages_df, default="https://example.com")
    saved_domain = st.session_state.get("gsc_target_domain", auto_domain)
    brand_tokens = extract_brand_tokens(saved_domain)

    # Candidate pages sorted by traffic
    candidate_urls = []
    if not pages_df.empty and 'page' in pages_df.columns:
        p_sorted = pages_df.sort_values('impressions', ascending=False) if 'impressions' in pages_df.columns else pages_df
        candidate_urls = [str(u).strip() for u in p_sorted['page'].dropna().unique() if str(u).strip()]

    # Map queries to target pages
    queries_df = map_queries_to_pages(queries_df, pages_df, base_domain=saved_domain)

    # Apply any manual overrides saved in session state
    if "gsc_manual_page_overrides" in st.session_state:
        for q_ov, p_ov in st.session_state["gsc_manual_page_overrides"].items():
            mask = queries_df['query'].astype(str).str.strip().str.lower() == q_ov.strip().lower()
            if mask.any():
                queries_df.loc[mask, 'target_page'] = p_ov
                queries_df.loc[mask, 'target_type'] = "🏠 Homepage (Brand)" if p_ov.rstrip('/') == saved_domain.rstrip('/') else "Custom Override"

    # 3. High-Level KPI Summary Cards
    total_clicks = int(queries_df['clicks'].sum()) if 'clicks' in queries_df.columns else 0
    total_impressions = int(queries_df['impressions'].sum()) if 'impressions' in queries_df.columns else 0
    avg_ctr = (total_clicks / total_impressions * 100) if total_impressions > 0 else 0.0
    avg_position = queries_df['position'].mean() if 'position' in queries_df.columns and len(queries_df) > 0 else 0.0
    total_queries = len(queries_df)
    total_pages = len(pages_df) if not pages_df.empty else 0

    st.markdown(f"""
    <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 14px; margin-bottom: 1.5rem;">
        <div style="background: rgba(15, 23, 42, 0.9); border: 1px solid rgba(56, 189, 248, 0.3); border-radius: 12px; padding: 1.2rem; text-align: center;">
            <div style="color: #94A3B8; font-size: 0.8rem; font-weight: 700; text-transform: uppercase; margin-bottom: 4px;">Total Clicks</div>
            <div style="font-size: 1.9rem; font-weight: 900; color: #38BDF8;">{total_clicks:,}</div>
            <div style="font-size: 0.74rem; color: #64748B; margin-top: 4px;">Organic Search Traffic</div>
        </div>
        <div style="background: rgba(15, 23, 42, 0.9); border: 1px solid rgba(139, 92, 246, 0.3); border-radius: 12px; padding: 1.2rem; text-align: center;">
            <div style="color: #94A3B8; font-size: 0.8rem; font-weight: 700; text-transform: uppercase; margin-bottom: 4px;">Total Impressions</div>
            <div style="font-size: 1.9rem; font-weight: 900; color: #A78BFA;">{total_impressions:,}</div>
            <div style="font-size: 0.74rem; color: #64748B; margin-top: 4px;">Total Search Visibility</div>
        </div>
        <div style="background: rgba(15, 23, 42, 0.9); border: 1px solid rgba(16, 185, 129, 0.3); border-radius: 12px; padding: 1.2rem; text-align: center;">
            <div style="color: #94A3B8; font-size: 0.8rem; font-weight: 700; text-transform: uppercase; margin-bottom: 4px;">Average CTR</div>
            <div style="font-size: 1.9rem; font-weight: 900; color: #34D399;">{avg_ctr:.2f}%</div>
            <div style="font-size: 0.74rem; color: #64748B; margin-top: 4px;">Click-Through Ratio</div>
        </div>
        <div style="background: rgba(15, 23, 42, 0.9); border: 1px solid rgba(245, 158, 11, 0.3); border-radius: 12px; padding: 1.2rem; text-align: center;">
            <div style="color: #94A3B8; font-size: 0.8rem; font-weight: 700; text-transform: uppercase; margin-bottom: 4px;">Average Position</div>
            <div style="font-size: 1.9rem; font-weight: 900; color: #FBBF24;">{avg_position:.1f}</div>
            <div style="font-size: 0.74rem; color: #64748B; margin-top: 4px;">Mean SERP Rank</div>
        </div>
        <div style="background: rgba(15, 23, 42, 0.9); border: 1px solid rgba(236, 72, 153, 0.3); border-radius: 12px; padding: 1.2rem; text-align: center;">
            <div style="color: #94A3B8; font-size: 0.8rem; font-weight: 700; text-transform: uppercase; margin-bottom: 4px;">Indexed Queries</div>
            <div style="font-size: 1.9rem; font-weight: 900; color: #F472B6;">{total_queries:,}</div>
            <div style="font-size: 0.74rem; color: #64748B; margin-top: 4px;">Unique Keyword Terms</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # 4. Target Domain & Page Mapping Bar
    with st.expander("🎯 Target Domain & Brand Keywords Settings", expanded=False):
        c_dom1, c_dom2 = st.columns([3, 1])
        with c_dom1:
            target_domain_input = st.text_input(
                "Website Base Domain for Target URLs:",
                value=saved_domain,
                placeholder="https://example.com",
                help="Domain used to map brand queries directly to Homepage and generate target URLs."
            )
            if target_domain_input != saved_domain:
                st.session_state["gsc_target_domain"] = target_domain_input
                st.rerun()
        with c_dom2:
            brand_count = len(queries_df[queries_df['target_type'].str.contains('Homepage', na=False)]) if 'target_type' in queries_df.columns else 0
            existing_count = len(queries_df[queries_df['target_type'] == 'Existing Page']) if 'target_type' in queries_df.columns else 0
            new_count = len(queries_df[queries_df['target_type'].str.contains('Create New', na=False)]) if 'target_type' in queries_df.columns else 0
            st.markdown("<div style='margin-top: 26px;'></div>", unsafe_allow_html=True)
            st.caption(f"🏠 **{brand_count}** Brand ➔ Homepage\n✔ **{existing_count}** Existing Pages\n💡 **{new_count}** Recommended New Pages")

    # 5. AI Configuration Bar (Matching user screenshot)
    st.markdown("""
    <div style="background: rgba(30, 41, 59, 0.6); border: 1px solid rgba(139, 92, 246, 0.3); border-radius: 14px; padding: 1.2rem 1.4rem; margin-bottom: 1.5rem;">
        <div style="display: flex; align-items: center; gap: 8px; margin-bottom: 12px;">
            <span style="font-size: 1.1rem;">🤖</span>
            <span style="font-weight: 800; font-size: 1.05rem; color: #F8FAFC;">AI SEO Optimization & Complete File Generator</span>
            <span style="font-size: 0.75rem; background: rgba(139, 92, 246, 0.2); color: #C4B5FD; padding: 2px 8px; border-radius: 4px; font-weight: 700;">GEMINI / OPENAI / CLAUDE</span>
        </div>
        <div style="font-size: 0.85rem; color: #94A3B8; margin-bottom: 14px; line-height: 1.5;">
            Target each query to the exact page (Brand ➔ Homepage), then generate <b>CTR-Optimized Title Tags</b>, <b>H1 Headings</b>, <b>Meta Descriptions</b>, and <b>Action Plans</b>, and download a complete strategy file.
        </div>
    """, unsafe_allow_html=True)

    c_prov, c_model, c_key = st.columns([1.2, 1.4, 2.0])

    with c_prov:
        ai_provider = st.selectbox(
            "AI Provider",
            options=["Google Gemini", "ChatGPT (OpenAI)", "Claude (Anthropic)"],
            index=0,
            key="gsc_ai_provider"
        )

    with c_model:
        if ai_provider == "Google Gemini":
            model_options = [
                "gemini-3.8-pro-preview",
                "gemini-3.8-flash",
                "gemini-3.7-pro",
                "gemini-3.7-flash",
                "gemini-2.5-flash",
                "gemini-2.5-pro",
                "gemini-2.0-flash",
                "Custom Model"
            ]
        elif ai_provider == "ChatGPT (OpenAI)":
            model_options = [
                "gpt-4o",
                "gpt-4o-mini",
                "gpt-4.1",
                "o3-mini",
                "Custom Model"
            ]
        else:
            model_options = [
                "claude-3-7-sonnet-20250219",
                "claude-3-5-sonnet-20241022",
                "claude-3-5-haiku-20241022",
                "Custom Model"
            ]

        selected_model_choice = st.selectbox(
            "AI Model",
            options=model_options,
            index=0,
            key="gsc_ai_model_choice"
        )
        if selected_model_choice == "Custom Model":
            selected_model = st.text_input(
                "Custom Model Name",
                placeholder="e.g. gemini-3.8-pro-preview",
                key="gsc_custom_model_input"
            ).strip()
        else:
            selected_model = selected_model_choice

    with c_key:
        if ai_provider == "Google Gemini":
            key_label = "Google Gemini API Key (Free)"
            key_placeholder = "Paste Gemini API key (AIzaSy...)"
            help_url = "https://aistudio.google.com/apikey"
            help_text = "Get free key from Google AI Studio →"
        elif ai_provider == "ChatGPT (OpenAI)":
            key_label = "OpenAI API Key"
            key_placeholder = "sk-..."
            help_url = "https://platform.openai.com/api-keys"
            help_text = "Get OpenAI API key →"
        else:
            key_label = "Anthropic API Key"
            key_placeholder = "sk-ant-..."
            help_url = "https://console.anthropic.com/settings/keys"
            help_text = "Get Anthropic key →"

        saved_key = st.session_state.get(f"api_key_{ai_provider}", "")
        api_key = st.text_input(
            key_label,
            value=saved_key,
            type="password",
            placeholder=key_placeholder,
            key=f"gsc_input_key_{ai_provider}"
        )
        if api_key:
            st.session_state[f"api_key_{ai_provider}"] = api_key

        st.markdown(
            f'<div style="font-size:0.75rem; margin-top:2px;"><a href="{help_url}" target="_blank" style="color:#38BDF8; text-decoration:none;">{help_text}</a></div>',
            unsafe_allow_html=True
        )

    c_gen1, c_gen2, c_gen3 = st.columns([2.0, 1.2, 1.2])
    with c_gen1:
        gen_scope = st.selectbox(
            "Scope for AI Optimization:",
            options=[
                "⚡ Priority Opportunities (Quick Wins + Low-Hanging CTR + Questions)",
                "🍎 Low-Hanging Fruits Only (Top 3 Pos, Low CTR)",
                "⚡ Page-2 Quick Wins Only (Striking Distance)",
                "❓ Questions & Informational Queries Only",
                "📊 Top 50 High-Volume Queries"
            ],
            index=0,
            key="gsc_gen_scope"
        )
    with c_gen2:
        max_ai_items = st.selectbox(
            "Max Queries to Optimize:",
            options=[15, 30, 50, 100],
            index=1,
            key="gsc_max_ai_items"
        )
    with c_gen3:
        st.markdown("<div style='margin-top: 26px;'></div>", unsafe_allow_html=True)
        btn_generate_ai = st.button("🚀 Generate AI Action Plan", use_container_width=True, type="primary")

    st.markdown("</div>", unsafe_allow_html=True)

    # 6. Global Analysis Tuning Controls
    with st.expander("⚙️ Analysis Parameters & Sensitivity Thresholds", expanded=False):
        c_p1, c_p2, c_p3 = st.columns(3)
        with c_p1:
            min_impressions = st.slider(
                "Minimum Impressions Filter:",
                min_value=1,
                max_value=10000,
                value=500 if total_impressions > 50000 else 20,
                step=10,
                help="Filters out low volume queries so you only focus on meaningful search traffic."
            )
            top_queries_count = st.slider("Top Queries to Display:", 5, 100, 25, step=5)
        with c_p2:
            min_pos_qw = st.slider("Quick Wins Min Position (Page 2 start):", 5, 20, 11, step=1)
            max_pos_qw = st.slider("Quick Wins Max Position (Page 2 end):", 15, 30, 20, step=1)
        with c_p3:
            max_pos_opp = st.slider("Content Gap Max Position Threshold:", 30, 100, 60, step=5)
            low_hanging_ctr_thresh = st.slider("Low-Hanging Fruit Max CTR Threshold:", 0.05, 0.40, 0.20, step=0.01, format="%.2f")

    def format_gsc_df_for_display(df: pd.DataFrame) -> pd.DataFrame:
        d = df.copy()
        if 'ctr' in d.columns:
            d['ctr'] = d['ctr'].apply(lambda x: f"{float(x)*100:.2f}%" if pd.notnull(x) else "0.00%")
        if 'position' in d.columns:
            d['position'] = d['position'].apply(lambda x: f"{float(x):.1f}" if pd.notnull(x) else "0.0")
        if 'clicks' in d.columns:
            d['clicks'] = d['clicks'].fillna(0).astype(int)
        if 'impressions' in d.columns:
            d['impressions'] = d['impressions'].fillna(0).astype(int)
        return d

    # Handle AI Generation trigger
    if btn_generate_ai:
        with st.spinner(f"✨ Reviewing site pages and generating AI recommendations using {selected_model}..."):
            candidates = []
            if "Low-Hanging Fruits Only" in gen_scope:
                lh_mask = (queries_df['position'] <= 3.5) & (queries_df['ctr'] < low_hanging_ctr_thresh) & (queries_df['impressions'] >= 5)
                sub = queries_df[lh_mask].sort_values(['impressions', 'ctr'], ascending=[False, True]).head(max_ai_items)
                for _, r in sub.iterrows():
                    candidates.append({**r.to_dict(), 'category': 'Low-Hanging Fruit'})
            elif "Page-2 Quick Wins Only" in gen_scope:
                qw_mask = (queries_df['position'] >= min_pos_qw) & (queries_df['position'] <= max_pos_qw) & (queries_df['impressions'] >= min_impressions)
                sub = queries_df[qw_mask].sort_values(['impressions', 'position'], ascending=[False, True]).head(max_ai_items)
                for _, r in sub.iterrows():
                    candidates.append({**r.to_dict(), 'category': 'Quick Win'})
            elif "Questions" in gen_scope:
                q_regex = r'\b(what|how|why|when|where|which|who|can|is|are|should|does|do)\b'
                quest_mask = queries_df['query'].astype(str).str.lower().str.contains(q_regex, regex=True, na=False)
                sub = queries_df[quest_mask].sort_values('impressions', ascending=False).head(max_ai_items)
                for _, r in sub.iterrows():
                    candidates.append({**r.to_dict(), 'category': 'Question'})
            elif "Top 50" in gen_scope:
                sub = queries_df.nlargest(min(max_ai_items, len(queries_df)), 'clicks')
                for _, r in sub.iterrows():
                    candidates.append({**r.to_dict(), 'category': 'Top Performer'})
            else:
                # Priority blend
                lh_mask = (queries_df['position'] <= 3.5) & (queries_df['ctr'] < low_hanging_ctr_thresh) & (queries_df['impressions'] >= 5)
                sub_lh = queries_df[lh_mask].sort_values('impressions', ascending=False).head(max_ai_items // 2)
                for _, r in sub_lh.iterrows():
                    candidates.append({**r.to_dict(), 'category': 'Low-Hanging Fruit'})

                qw_mask = (queries_df['position'] >= min_pos_qw) & (queries_df['position'] <= max_pos_qw)
                sub_qw = queries_df[qw_mask].sort_values('impressions', ascending=False).head(max_ai_items // 2)
                for _, r in sub_qw.iterrows():
                    candidates.append({**r.to_dict(), 'category': 'Quick Win'})

                q_regex = r'\b(what|how|why|when|where|which|who|can|is|are|should|does|do)\b'
                quest_mask = queries_df['query'].astype(str).str.lower().str.contains(q_regex, regex=True, na=False)
                sub_quest = queries_df[quest_mask].sort_values('impressions', ascending=False).head(10)
                for _, r in sub_quest.iterrows():
                    candidates.append({**r.to_dict(), 'category': 'Question'})

            unique_candidates = []
            seen_q = set()
            for c in candidates:
                q_txt = str(c.get('query', '')).strip().lower()
                if q_txt and q_txt not in seen_q:
                    seen_q.add(q_txt)
                    unique_candidates.append(c)

            unique_candidates = unique_candidates[:max_ai_items]

            if not unique_candidates:
                sub = queries_df.nlargest(min(max_ai_items, len(queries_df)), 'impressions')
                for _, r in sub.iterrows():
                    unique_candidates.append({**r.to_dict(), 'category': 'Top Performer'})

            ai_results = generate_ai_seo_action_plan(
                unique_candidates,
                api_key=api_key,
                provider=ai_provider,
                model=selected_model,
                candidate_pages=candidate_urls,
                base_domain=saved_domain,
                brand_tokens=brand_tokens
            )
            st.session_state["gsc_ai_action_plan"] = ai_results
            st.toast("✅ AI Action Plan & Metadata successfully generated!", icon="🎉")

    # Display / Download AI Action Plan Banner if results are available
    if "gsc_ai_action_plan" in st.session_state and st.session_state["gsc_ai_action_plan"]:
        ai_df = pd.DataFrame(st.session_state["gsc_ai_action_plan"])
        st.markdown(f"""
        <div style="background: rgba(16, 185, 129, 0.1); border: 1px solid rgba(16, 185, 129, 0.4); border-radius: 12px; padding: 1.1rem 1.4rem; margin-bottom: 1.2rem;">
            <div style="display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 10px;">
                <div>
                    <div style="font-weight: 800; font-size: 1.05rem; color: #34D399; margin-bottom: 2px;">
                        ✨ Complete SEO Action Plan Ready ({len(ai_df)} Queries Optimized)
                    </div>
                    <div style="font-size: 0.84rem; color: #94A3B8;">
                        Includes Target URLs (Brand ➔ Homepage), Clicks, Impressions, CTR, Positions, AI Title Tags, H1 Headings, Meta Descriptions, and Action Plans.
                    </div>
                </div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        col_dl_xl, col_dl_csv = st.columns(2)
        with col_dl_xl:
            excel_bytes = export_action_plan_to_excel(ai_df)
            st.download_button(
                "📥 Download Complete SEO Action Plan (.XLSX)",
                data=excel_bytes,
                file_name="complete_seo_action_plan.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="btn_dl_ai_excel",
                use_container_width=True
            )
        with col_dl_csv:
            csv_bytes = ai_df.to_csv(index=False).encode('utf-8')
            st.download_button(
                "📥 Download Complete SEO Action Plan (.CSV)",
                data=csv_bytes,
                file_name="complete_seo_action_plan.csv",
                mime="text/csv",
                key="btn_dl_ai_csv",
                use_container_width=True
            )

    # 7. Core Analytical Modules Tabs
    tab_ai, tab_qw, tab_lh, tab_pot, tab_top, tab_quest, tab_opp, tab_cloud, tab_geo_page = st.tabs([
        "🤖 AI SEO Action Plan",
        "⚡ Quick Wins (Pos 11-20)",
        "🍎 Low-Hanging Fruits (CTR)",
        "📈 Traffic Potential Estimator",
        "📊 Top Performing Queries",
        "❓ Questions & Informational",
        "🎯 Keyword Gaps (Pos > 60)",
        "☁️ Keyword Word Cloud",
        "🌍 Countries & Pages"
    ])

    query_cols = [c for c in ['query', 'target_page', 'target_type', 'clicks', 'impressions', 'ctr', 'position'] if c in queries_df.columns]

    # TAB 0: AI ACTION PLAN
    with tab_ai:
        if "gsc_ai_action_plan" in st.session_state and st.session_state["gsc_ai_action_plan"]:
            plan_df = pd.DataFrame(st.session_state["gsc_ai_action_plan"])
            plan_df.index = range(1, len(plan_df) + 1)

            st.markdown("#### 🎯 AI Recommended Target Pages & Content Metadata")
            disp_plan = plan_df.copy()
            if 'ctr' in disp_plan.columns:
                disp_plan['ctr'] = disp_plan['ctr'].apply(lambda x: f"{float(x)*100:.2f}%" if pd.notnull(x) else "0.00%")
            if 'position' in disp_plan.columns:
                disp_plan['position'] = disp_plan['position'].apply(lambda x: f"{float(x):.1f}" if pd.notnull(x) else "0.0")

            display_cols = [c for c in ['query', 'target_page', 'target_type', 'ai_title', 'ai_h1', 'ai_meta_description', 'action_plan', 'clicks', 'impressions', 'ctr', 'position'] if c in disp_plan.columns]
            st.dataframe(disp_plan[display_cols], use_container_width=True, height=420)

            # Interactive Quick Target Page Override
            with st.expander("✏️ Quick Target Page Customizer & Manual Switcher", expanded=False):
                st.caption("Need to adjust the target page for any query? Select below to assign to Homepage or any specific URL.")
                c_ov1, c_ov2, c_ov3 = st.columns([2, 3, 1.2])
                with c_ov1:
                    query_list = [str(r.get('query')) for r in st.session_state["gsc_ai_action_plan"]]
                    selected_q_ov = st.selectbox("Select Query to Adjust:", options=query_list, key="sel_q_ov")
                with c_ov2:
                    current_ov_val = next((r.get('target_page', '') for r in st.session_state["gsc_ai_action_plan"] if str(r.get('query')) == selected_q_ov), "")
                    new_ov_val = st.text_input("Target URL:", value=current_ov_val, key="input_target_ov")
                with c_ov3:
                    st.markdown("<div style='margin-top: 26px;'></div>", unsafe_allow_html=True)
                    btn_save_ov = st.button("Apply URL", key="btn_save_target_ov", use_container_width=True)

                c_quick1, c_quick2 = st.columns(2)
                with c_quick1:
                    if st.button(f"🏠 Set '{selected_q_ov}' directly to Homepage", key="btn_quick_home", use_container_width=True):
                        new_ov_val = saved_domain.rstrip('/') + '/'
                        btn_save_ov = True
                with c_quick2:
                    if candidate_urls:
                        picked_existing = st.selectbox("Or Pick From Website's Top Pages:", options=["-- Select Page --"] + candidate_urls[:25], key="sel_pick_existing")
                        if picked_existing != "-- Select Page --":
                            new_ov_val = picked_existing
                            btn_save_ov = True

                if btn_save_ov:
                    if "gsc_manual_page_overrides" not in st.session_state:
                        st.session_state["gsc_manual_page_overrides"] = {}
                    st.session_state["gsc_manual_page_overrides"][selected_q_ov] = new_ov_val

                    # Update active plan
                    for r in st.session_state["gsc_ai_action_plan"]:
                        if str(r.get('query')) == selected_q_ov:
                            r['target_page'] = new_ov_val
                            r['target_type'] = "🏠 Homepage (Brand)" if new_ov_val.rstrip('/') == saved_domain.rstrip('/') else "Custom Override"

                    st.toast(f"✅ Target Page updated for '{selected_q_ov}'!", icon="🚀")
                    st.rerun()
        else:
            st.markdown("""
            <div style="background: rgba(15, 23, 42, 0.7); border: 1px dashed rgba(139, 92, 246, 0.4); border-radius: 14px; padding: 2rem; text-align: center;">
                <div style="font-size: 2.2rem; margin-bottom: 0.5rem;">🤖</div>
                <h4 style="color: #F8FAFC; margin-bottom: 0.4rem;">No AI Action Plan Generated Yet</h4>
                <p style="color: #94A3B8; font-size: 0.9rem; max-width: 580px; margin: 0 auto 1.2rem; line-height: 1.6;">
                    Click the <b>'🚀 Generate AI Action Plan'</b> button in the AI configuration box above to review all site pages and create high-converting Page Titles, H1 tags, Meta Descriptions, and action plans.
                </p>
            </div>
            """, unsafe_allow_html=True)

    # TAB 1: QUICK WINS
    with tab_qw:
        st.markdown(f"""
        <div style="background: rgba(16, 185, 129, 0.08); border: 1px solid rgba(16, 185, 129, 0.3); border-radius: 12px; padding: 1.1rem 1.4rem; margin-bottom: 1.2rem;">
            <div style="font-weight: 800; font-size: 1.05rem; color: #10B981; margin-bottom: 4px;">⚡ Page-2 Striking Distance Keywords (Position {min_pos_qw} - {max_pos_qw})</div>
            <div style="font-size: 0.85rem; color: #94A3B8; line-height: 1.5;">
                These keywords rank on <b>Page 2 of Google</b> with at least <b>{min_impressions} impressions</b>. Target these queries on the listed target URLs with small on-page optimizations (adding keywords to H2s, internal links, and content refreshes) to push them into Page 1.
            </div>
        </div>
        """, unsafe_allow_html=True)

        filter_mask = (
            (queries_df['position'] >= min_pos_qw) &
            (queries_df['position'] <= max_pos_qw) &
            (queries_df['impressions'] >= min_impressions)
        )
        sort_cols = [c for c in ['impressions', 'position'] if c in queries_df.columns]
        if sort_cols:
            quick_wins = queries_df[filter_mask][query_cols].sort_values(sort_cols, ascending=[False, True]).reset_index(drop=True)
        else:
            quick_wins = queries_df[filter_mask][query_cols].reset_index(drop=True)

        if len(quick_wins) > 0:
            quick_wins.index = range(1, len(quick_wins) + 1)
            qw_disp = format_gsc_df_for_display(quick_wins)
            st.dataframe(qw_disp, use_container_width=True, height=380)

            qw_csv = quick_wins.to_csv(index=False).encode('utf-8')
            st.download_button(
                "📥 Download Quick Wins Report (.CSV)",
                data=qw_csv,
                file_name="gsc_quick_wins_striking_distance.csv",
                mime="text/csv",
                key="btn_dl_qw"
            )
        else:
            st.info(f"No quick wins found with position between {min_pos_qw}-{max_pos_qw} and ≥ {min_impressions} impressions. Try lowering the impressions threshold in settings above.")

    # TAB 2: LOW-HANGING FRUITS
    with tab_lh:
        st.markdown(f"""
        <div style="background: rgba(245, 158, 11, 0.08); border: 1px solid rgba(245, 158, 11, 0.3); border-radius: 12px; padding: 1.1rem 1.4rem; margin-bottom: 1.2rem;">
            <div style="font-weight: 800; font-size: 1.05rem; color: #F59E0B; margin-bottom: 4px;">🍎 Low-Hanging Fruit Optimization (Top 3 Positions, CTR < {low_hanging_ctr_thresh*100:.0f}%)</div>
            <div style="font-size: 0.85rem; color: #94A3B8; line-height: 1.5;">
                These queries already hold prime real estate in the <b>Top 3 search results</b>, but searchers aren't clicking. By rewriting your <b>Page Titles and Meta Descriptions</b> on the target pages (Brand queries target Homepage) to be more compelling, adding numbers, power words, and answering search intent, you can double your clicks without building a single backlink.
            </div>
        </div>
        """, unsafe_allow_html=True)

        filter_mask = (
            (queries_df['position'] <= 3.5) &
            (queries_df['ctr'] < low_hanging_ctr_thresh) &
            (queries_df['impressions'] >= max(5, min_impressions // 2))
        )
        sort_cols = [c for c in ['impressions', 'ctr'] if c in queries_df.columns]
        if sort_cols:
            low_hanging = queries_df[filter_mask][query_cols].sort_values(sort_cols, ascending=[False, True]).reset_index(drop=True)
        else:
            low_hanging = queries_df[filter_mask][query_cols].reset_index(drop=True)

        if len(low_hanging) > 0:
            low_hanging.index = range(1, len(low_hanging) + 1)
            lh_disp = format_gsc_df_for_display(low_hanging)
            st.dataframe(lh_disp, use_container_width=True, height=380)

            lh_csv = low_hanging.to_csv(index=False).encode('utf-8')
            st.download_button(
                "📥 Download Low-Hanging Fruits Report (.CSV)",
                data=lh_csv,
                file_name="gsc_low_hanging_fruits.csv",
                mime="text/csv",
                key="btn_dl_lh"
            )
        else:
            st.info("No low-hanging fruit queries found with current thresholds. All top 3 positions have healthy CTR!")

    # TAB 3: TRAFFIC POTENTIAL ESTIMATION
    with tab_pot:
        st.markdown("""
        <div style="background: rgba(56, 189, 248, 0.08); border: 1px solid rgba(56, 189, 248, 0.3); border-radius: 12px; padding: 1.1rem 1.4rem; margin-bottom: 1.2rem;">
            <div style="font-weight: 800; font-size: 1.05rem; color: #38BDF8; margin-bottom: 4px;">📈 Algorithmic Traffic Growth Projection Engine</div>
            <div style="font-size: 0.85rem; color: #94A3B8; line-height: 1.5;">
                How many extra clicks would your site generate if sub-page-1 keywords reached the first page? This model benchmarks your domain's actual <b>Top 10 average CTR</b> and projects incremental clicks across lower-ranking keywords.
            </div>
        </div>
        """, unsafe_allow_html=True)

        top10_mask = queries_df['position'] <= 10.0
        top10_clicks = queries_df[top10_mask]['clicks'].sum() if 'clicks' in queries_df.columns else 0
        top10_impressions = queries_df[top10_mask]['impressions'].sum() if 'impressions' in queries_df.columns else 0
        avg_ctr_top10 = (top10_clicks / top10_impressions) if top10_impressions > 0 else 0.035

        calc_df = queries_df.copy()
        if 'impressions' in calc_df.columns and 'clicks' in calc_df.columns and 'position' in calc_df.columns:
            calc_df['potential_clicks'] = calc_df.apply(
                lambda r: r['clicks'] if r['position'] <= 10.0 else (r['impressions'] * avg_ctr_top10),
                axis=1
            )
            calc_df['click_gain'] = calc_df['potential_clicks'] - calc_df['clicks']
        else:
            calc_df['potential_clicks'] = 0
            calc_df['click_gain'] = 0

        current_clicks_sum = calc_df['clicks'].sum() if 'clicks' in calc_df.columns else 0
        potential_clicks_sum = calc_df['potential_clicks'].sum()
        gain_clicks_sum = max(0, potential_clicks_sum - current_clicks_sum)
        gain_pct = (gain_clicks_sum / current_clicks_sum * 100) if current_clicks_sum > 0 else 0

        pot_col1, pot_col2, pot_col3, pot_col4 = st.columns(4)
        with pot_col1:
            st.metric("Top-10 Benchmark CTR", f"{avg_ctr_top10*100:.2f}%")
        with pot_col2:
            st.metric("Current Total Clicks", f"{int(current_clicks_sum):,}")
        with pot_col3:
            st.metric("Estimated Potential Clicks", f"{int(potential_clicks_sum):,}")
        with pot_col4:
            st.metric("Projected Growth", f"+{gain_pct:.1f}%", f"+{int(gain_clicks_sum):,} extra clicks")

        pot_target_cols = [c for c in ['query', 'target_page', 'position', 'clicks', 'potential_clicks', 'click_gain', 'impressions'] if c in calc_df.columns]
        potential_table = calc_df[
            (calc_df['click_gain'] > 1) &
            (calc_df['position'] > 10.0)
        ][pot_target_cols]

        if 'click_gain' in potential_table.columns:
            potential_table = potential_table.sort_values('click_gain', ascending=False)
        potential_table = potential_table.head(30).reset_index(drop=True)

        if len(potential_table) > 0:
            potential_table.index = range(1, len(potential_table) + 1)
            disp_pot = potential_table.copy()
            if 'clicks' in disp_pot.columns: disp_pot['clicks'] = disp_pot['clicks'].astype(int)
            if 'potential_clicks' in disp_pot.columns: disp_pot['potential_clicks'] = disp_pot['potential_clicks'].astype(int)
            if 'click_gain' in disp_pot.columns: disp_pot['click_gain'] = disp_pot['click_gain'].astype(int)
            if 'impressions' in disp_pot.columns: disp_pot['impressions'] = disp_pot['impressions'].astype(int)
            if 'position' in disp_pot.columns: disp_pot['position'] = disp_pot['position'].apply(lambda x: f"{float(x):.1f}")

            st.markdown("#### 🎯 Top Keywords Ranked by Click Gain Opportunity:")
            st.dataframe(disp_pot, use_container_width=True, height=350)

            pot_csv = potential_table.to_csv(index=False).encode('utf-8')
            st.download_button(
                "📥 Download Traffic Potential Report (.CSV)",
                data=pot_csv,
                file_name="gsc_traffic_potential_opportunities.csv",
                mime="text/csv",
                key="btn_dl_pot"
            )
        else:
            st.info("No queries found with significant gain potential based on current impressions.")

    # TAB 4: TOP PERFORMING QUERIES
    with tab_top:
        st.markdown(f"#### 📊 Top {top_queries_count} Queries Driving Highest Clicks")
        sort_col = 'clicks' if 'clicks' in queries_df.columns else ('impressions' if 'impressions' in queries_df.columns else queries_df.columns[0])
        top_queries = queries_df.nlargest(min(top_queries_count, len(queries_df)), sort_col)[query_cols].reset_index(drop=True)
        top_queries.index = range(1, len(top_queries) + 1)

        disp_top = format_gsc_df_for_display(top_queries)
        st.dataframe(disp_top, use_container_width=True, height=380)

        if len(top_queries) > 0 and 'clicks' in top_queries.columns:
            fig_top = px.bar(
                top_queries.head(15).sort_values('clicks', ascending=True),
                x='clicks',
                y='query',
                orientation='h',
                title="Top 15 Queries by Organic Clicks",
                color='clicks',
                color_continuous_scale='Purples',
                template="plotly_dark"
            )
            fig_top.update_layout(
                paper_bgcolor='rgba(15, 23, 42, 0.9)',
                plot_bgcolor='rgba(15, 23, 42, 0.9)',
                margin=dict(l=20, r=20, t=40, b=20)
            )
            st.plotly_chart(fig_top, use_container_width=True)

    # TAB 5: QUESTION QUERIES
    with tab_quest:
        st.markdown("""
        <div style="background: rgba(139, 92, 246, 0.08); border: 1px solid rgba(139, 92, 246, 0.3); border-radius: 12px; padding: 1.1rem 1.4rem; margin-bottom: 1.2rem;">
            <div style="font-weight: 800; font-size: 1.05rem; color: #A78BFA; margin-bottom: 4px;">❓ Informational & Question Search Queries</div>
            <div style="font-size: 0.85rem; color: #94A3B8; line-height: 1.5;">
                These queries contain question modifiers (<i>what, how, why, when, where, which, can, is, are</i>). Target them with dedicated FAQ sections, Featured Snippets, and educational guides on the mapped URLs.
            </div>
        </div>
        """, unsafe_allow_html=True)

        question_regex = r'\b(what|how|why|when|where|which|who|can|is|are|should|does|do)\b'
        question_queries = queries_df[
            queries_df['query'].astype(str).str.lower().str.contains(question_regex, regex=True, na=False)
        ][query_cols]

        if 'impressions' in question_queries.columns:
            question_queries = question_queries.sort_values('impressions', ascending=False)
        question_queries = question_queries.reset_index(drop=True)

        if len(question_queries) > 0:
            question_queries.index = range(1, len(question_queries) + 1)
            disp_quest = format_gsc_df_for_display(question_queries)
            st.dataframe(disp_quest, use_container_width=True, height=380)

            q_csv = question_queries.to_csv(index=False).encode('utf-8')
            st.download_button(
                "📥 Download Question Queries (.CSV)",
                data=q_csv,
                file_name="gsc_question_queries.csv",
                mime="text/csv",
                key="btn_dl_quest"
            )
        else:
            st.info("No question queries found matching question terms (what, how, why, etc.).")

    # TAB 6: KEYWORD OPPORTUNITIES
    with tab_opp:
        st.markdown(f"""
        <div style="background: rgba(239, 68, 68, 0.08); border: 1px solid rgba(239, 68, 68, 0.3); border-radius: 12px; padding: 1.1rem 1.4rem; margin-bottom: 1.2rem;">
            <div style="font-weight: 800; font-size: 1.05rem; color: #EF4444; margin-bottom: 4px;">🎯 Content Gap Opportunities (Position > {max_pos_opp})</div>
            <div style="font-size: 0.85rem; color: #94A3B8; line-height: 1.5;">
                These keywords rank beyond <b>Position {max_pos_opp}</b> but already generate at least <b>{min_impressions} impressions</b>. Search engines associate your site with the topic, but you need a dedicated landing page. Target these with the recommended URLs below.
            </div>
        </div>
        """, unsafe_allow_html=True)

        filter_mask = (
            (queries_df['position'] > max_pos_opp) &
            (queries_df['impressions'] >= min_impressions)
        )
        sort_cols = [c for c in ['impressions', 'position'] if c in queries_df.columns]
        if sort_cols:
            opps = queries_df[filter_mask][query_cols].sort_values(sort_cols, ascending=[False, True]).reset_index(drop=True)
        else:
            opps = queries_df[filter_mask][query_cols].reset_index(drop=True)

        if len(opps) > 0:
            opps.index = range(1, len(opps) + 1)
            disp_opp = format_gsc_df_for_display(opps)
            st.dataframe(disp_opp, use_container_width=True, height=380)

            opp_csv = opps.to_csv(index=False).encode('utf-8')
            st.download_button(
                "📥 Download Content Gap Opportunities (.CSV)",
                data=opp_csv,
                file_name="gsc_content_gap_opportunities.csv",
                mime="text/csv",
                key="btn_dl_opp"
            )
        else:
            st.info(f"No keyword opportunities found beyond position {max_pos_opp} with ≥ {min_impressions} impressions.")

    # TAB 7: KEYWORD WORD CLOUD
    with tab_cloud:
        st.markdown("#### ☁️ Keyword Semantic Themes & Word Cloud")
        st.write("Visualizes the most frequently occurring terms across your top queries weighted by search impressions.")

        try:
            if 'impressions' in queries_df.columns:
                top_for_cloud = queries_df.nlargest(min(150, len(queries_df)), 'impressions')['query'].astype(str).tolist()
            else:
                top_for_cloud = queries_df['query'].head(150).astype(str).tolist()
            text = ' '.join(top_for_cloud)
            if len(text.strip()) > 5:
                wc = WordCloud(
                    width=900,
                    height=450,
                    background_color='#0F172A',
                    colormap='plasma',
                    max_words=100
                ).generate(text)

                fig_wc, ax = plt.subplots(figsize=(12, 6), facecolor='#0F172A')
                ax.imshow(wc, interpolation='bilinear')
                ax.axis('off')
                st.pyplot(fig_wc)
                plt.close(fig_wc)
            else:
                st.info("Not enough text to generate a word cloud.")
        except Exception as e:
            st.error(f"Could not generate word cloud: {str(e)}")

    # TAB 8: COUNTRIES & PAGES
    with tab_geo_page:
        col_g1, col_g2 = st.columns(2)

        if not countries_df.empty:
            c_name_col = 'country' if 'country' in countries_df.columns else countries_df.columns[0]
            with col_g1:
                st.markdown("#### 🌍 Top Countries by Clicks")
                sort_col = 'clicks' if 'clicks' in countries_df.columns else countries_df.columns[min(1, len(countries_df.columns)-1)]
                top_countries = countries_df.sort_values(sort_col, ascending=False).head(10)
                fig_c = px.bar(
                    top_countries.sort_values(sort_col, ascending=True),
                    y=c_name_col,
                    x=sort_col,
                    orientation='h',
                    color=sort_col,
                    color_continuous_scale='Blues',
                    template="plotly_dark"
                )
                fig_c.update_layout(paper_bgcolor='rgba(15, 23, 42, 0.9)', plot_bgcolor='rgba(15, 23, 42, 0.9)', margin=dict(l=20, r=20, t=30, b=20))
                st.plotly_chart(fig_c, use_container_width=True)

            with col_g2:
                st.markdown("#### 🌐 Top Countries by Impressions")
                sort_col_i = 'impressions' if 'impressions' in countries_df.columns else countries_df.columns[min(2, len(countries_df.columns)-1)]
                fig_ci = px.bar(
                    top_countries.sort_values(sort_col_i, ascending=True) if sort_col_i in top_countries.columns else top_countries,
                    y=c_name_col,
                    x=sort_col_i if sort_col_i in top_countries.columns else sort_col,
                    orientation='h',
                    color=sort_col_i if sort_col_i in top_countries.columns else sort_col,
                    color_continuous_scale='Greens',
                    template="plotly_dark"
                )
                fig_ci.update_layout(paper_bgcolor='rgba(15, 23, 42, 0.9)', plot_bgcolor='rgba(15, 23, 42, 0.9)', margin=dict(l=20, r=20, t=30, b=20))
                st.plotly_chart(fig_ci, use_container_width=True)

        if not pages_df.empty:
            st.markdown("<hr style='border-color: rgba(255,255,255,0.08); margin: 1.5rem 0;'>", unsafe_allow_html=True)
            st.markdown("#### 📄 Top Pages Driving Organic Traffic")
            p_cols = [c for c in ['page', 'clicks', 'impressions', 'ctr', 'position'] if c in pages_df.columns]
            sort_p = 'clicks' if 'clicks' in pages_df.columns else pages_df.columns[0]
            top_pages = pages_df.sort_values(sort_p, ascending=False).head(20).reset_index(drop=True)
            top_pages.index = range(1, len(top_pages) + 1)
            st.dataframe(format_gsc_df_for_display(top_pages[p_cols]), use_container_width=True, height=300)

            st.markdown("#### 📌 Underperforming Pages (High Impressions, Low CTR < 4%)")
            under_mask = (
                (pages_df['impressions'] >= min_impressions) &
                (pages_df['ctr'] < 0.04)
            )
            sort_under = 'impressions' if 'impressions' in pages_df.columns else pages_df.columns[0]
            under_pages = pages_df[under_mask].sort_values(sort_under, ascending=False).head(15).reset_index(drop=True)

            if len(under_pages) > 0:
                under_pages.index = range(1, len(under_pages) + 1)
                st.dataframe(format_gsc_df_for_display(under_pages[p_cols]), use_container_width=True, height=280)
            else:
                st.info("No pages found with high impressions and CTR < 4%.")
